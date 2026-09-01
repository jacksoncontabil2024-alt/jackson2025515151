"""Monta o modelo de slides dinâmico e gera o PPTX EDITÁVEL no padrão FELCONT.
Todos os elementos são caixas de texto, tabelas e gráficos nativos (editáveis)."""
import os
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE

from .indicators import brl, INSUF

# paleta oficial
INDIGO = RGBColor(0x32, 0x2F, 0x6A)
INDIGO2 = RGBColor(0x3E, 0x3A, 0x82)
TEAL = RGBColor(0x04, 0xB7, 0xAF)
TEAL_LT = RGBColor(0xC7, 0xEE, 0xEB)
AMBER = RGBColor(0xE8, 0xA1, 0x3A)
GREEN = RGBColor(0x57, 0xB1, 0x4A)
RED = RGBColor(0xD6, 0x45, 0x3F)
INK = RGBColor(0x26, 0x24, 0x45)
GRAY = RGBColor(0x6B, 0x6E, 0x86)
GRAY_LT = RGBColor(0xF2, 0xF3, 0xF8)
GRAY_BD = RGBColor(0xDE, 0xE0, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TITLE_FONT = "Roboto"
BODY_FONT = "Open Sans"
LOGO_PNG = "/app/assets_in/a2.png"
SW, SH = 13.333, 7.5


# ============================================================ modelo de slides
def build_slides(analysis):
    fin = analysis.get("financials") or {}
    ind = analysis.get("indicators") or {}
    diag = analysis.get("diagnosis") or {}
    meta = analysis.get("meta") or {}
    slides = []

    def add(stype, title, data, subtitle="", notes=""):
        slides.append({
            "id": str(uuid.uuid4())[:8], "type": stype, "title": title,
            "subtitle": subtitle, "visible": True, "notes": notes, "data": data,
        })

    add("capa", "Relatório Gerencial", {
        "client": meta.get("client_name") or analysis.get("client_name") or "",
        "period": meta.get("period_label") or analysis.get("period_label") or "",
        "responsavel": analysis.get("responsavel") or "",
    })

    cards = ind.get("cards") or []
    if cards:
        add("kpis", "Visão Executiva",
            {"items": [_kpi(c) for c in cards[:8]]},
            subtitle="Principais números do período")

    dre = fin.get("dre")
    if dre:
        add("resumo", "Resumo da Operação",
            {"items": [_kpi(c) for c in (ind.get("margens") or [])]},
            subtitle="Margens e rentabilidade")
        add("barras", "Receita e Resultado", _dre_bars(ind),
            subtitle="Comparativo das principais linhas da DRE")

    caixa = fin.get("caixa_mensal") or []
    if caixa:
        add("barras", "Disponibilidade de Caixa", {
            "categories": [c.get("mes") for c in caixa],
            "series": [{"name": "Caixa", "color": "indigo",
                        "values": [c.get("caixa") for c in caixa]}],
            "unit": "R$",
        }, subtitle="Evolução mensal do caixa")

    liq = ind.get("liquidez") or []
    if any(x.get("value") is not None for x in liq):
        add("kpis", "Liquidez", {"items": [_kpi(c) for c in liq]},
            subtitle="Índices de liquidez do período")

    folha = fin.get("folha")
    if folha:
        add("kpis", "Folha de Pagamento", {"items": _folha_kpis(folha)},
            subtitle="Composição e custo de pessoal")

    if dre:
        add("dre", "DRE Gerencial", {"dre": dre}, subtitle="Demonstração do Resultado — padrão FELCONT")
    bal = fin.get("balanco")
    if bal:
        add("balanco", "Balanço Patrimonial", {"balanco": bal},
            subtitle="Estrutura patrimonial — padrão FELCONT")

    pts = [d for d in (diag.get("diagnostico") or [])
           if d.get("tipo") in ("Ponto de Atenção", "Risco")]
    if pts:
        add("lista", "Pontos de Atenção", {"items": pts, "accent": "amber"})
    recs = diag.get("recomendacoes") or []
    if recs:
        add("lista", "Recomendações FELCONT",
            {"items": [{"titulo": r.get("titulo"), "texto": r.get("texto")} for r in recs],
             "accent": "teal"})

    add("conclusao", "Conclusão", {
        "resumo": diag.get("resumo_executivo") or "",
        "client": meta.get("client_name") or analysis.get("client_name") or "",
    })
    return slides


def _kpi(c):
    acc = "teal"
    if c.get("unit") == "%" and (c.get("value") or 0) < 0:
        acc = "red"
    if c.get("key") in ("resultado_liquido",) and (c.get("value") or 0) < 0:
        acc = "red"
    return {"label": c.get("label"), "value": c.get("display"), "accent": acc,
            "formula": c.get("formula"), "sources": c.get("sources")}


def _dre_bars(ind):
    cmp = ind.get("computed") or {}
    order = [("Receita Líquida", cmp.get("rol"), "teal"),
             ("Lucro Bruto", cmp.get("lucro_bruto"), "green"),
             ("Resultado Op.", cmp.get("res_op"), "indigo"),
             ("Resultado Líq.", cmp.get("res_liq"), "amber")]
    cats = [o[0] for o in order if o[1] is not None]
    vals = [o[1] for o in order if o[1] is not None]
    cols = [o[2] for o in order if o[1] is not None]
    return {"categories": cats, "series": [{"name": "R$", "values": vals, "colors": cols}],
            "unit": "R$", "percolor": True}


def _folha_kpis(folha):
    from .indicators import num, brl as _b
    items = []
    if folha.get("custo_total") is not None:
        items.append({"label": "Custo Total da Folha", "value": _b(num(folha["custo_total"])), "accent": "indigo"})
    for k, lab in [("proventos", "Proventos"), ("encargos", "Encargos Sociais"),
                   ("beneficios", "Benefícios"), ("provisoes", "Provisões")]:
        if folha.get(k) is not None:
            items.append({"label": lab, "value": _b(num(folha[k])), "accent": "teal"})
    if folha.get("num_colaboradores") is not None:
        items.append({"label": "Colaboradores", "value": str(int(num(folha["num_colaboradores"]))), "accent": "green"})
    return items or [{"label": "Folha", "value": INSUF, "accent": "amber"}]


# ============================================================ helpers de desenho
_ACC = {"teal": TEAL, "indigo": INDIGO, "amber": AMBER, "green": GREEN, "red": RED}


def _bg(s, color):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SW), Inches(SH))
    r.fill.solid(); r.fill.fore_color.rgb = color; r.line.fill.background(); r.shadow.inherit = False
    return r


def _rect(s, l, t, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0] = radius
        except Exception: pass
    return sp


def _txt(s, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0, sp_after=2):
    box = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{m}", 0)
    paras = runs if runs and isinstance(runs[0], list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing; p.space_after = Pt(sp_after); p.space_before = Pt(0)
        for (text, size, color, bold, font) in para:
            r = p.add_run(); r.text = text
            r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = font
    return box


def _header(s, title, subtitle, num, total):
    _rect(s, 0, 0, SW, 0.14, fill=TEAL)
    _txt(s, 0.62, 0.5, 10, 0.4, [[(subtitle.upper() if subtitle else "FELCONT REPORTS", 11, TEAL, True, BODY_FONT)]])
    _txt(s, 0.62, 0.84, 11.4, 0.9, [[(title, 27, INK, True, TITLE_FONT)]])
    _rect(s, 0.63, 1.6, 1.2, 0.06, fill=AMBER)
    _txt(s, SW - 1.6, 0.55, 1.0, 0.5, [[(f"{num:02d}", 20, GRAY_BD, True, TITLE_FONT)],
         [(f"/ {total:02d}", 10, GRAY_BD, False, BODY_FONT)]], align=PP_ALIGN.RIGHT)


def _footer(s, note):
    _rect(s, 0.62, SH - 0.5, 12.09, 0.012, fill=GRAY_BD)
    _txt(s, 0.62, SH - 0.42, 10.5, 0.32, [[(note, 8, GRAY, False, BODY_FONT)]])
    _txt(s, SW - 2.4, SH - 0.42, 1.78, 0.32, [[("FELCONT", 8.5, TEAL, True, BODY_FONT)]], align=PP_ALIGN.RIGHT)


# ============================================================ renderers
def _r_capa(s, sl, meta):
    _bg(s, INDIGO)
    _rect(s, 0, 0, SW, 0.14, fill=TEAL)
    if os.path.exists(LOGO_PNG):
        s.shapes.add_picture(LOGO_PNG, Inches(0.7), Inches(0.6), width=Inches(3.0))
    d = sl["data"]
    _txt(s, 0.75, 2.7, 9, 0.5, [[("RELATÓRIO GERENCIAL", 15, TEAL, True, BODY_FONT)]])
    _txt(s, 0.75, 3.2, 11.5, 1.6, [[(d.get("client") or "Cliente", 42, WHITE, True, TITLE_FONT)]], spacing=1.02)
    _rect(s, 0.78, 4.75, 1.5, 0.06, fill=AMBER)
    rows = []
    if d.get("period"):
        rows.append([("Período:  ", 14, GRAY_LT, False, BODY_FONT), (d["period"], 14, WHITE, True, BODY_FONT)])
    if d.get("responsavel"):
        rows.append([("Responsável:  ", 14, GRAY_LT, False, BODY_FONT), (d["responsavel"], 14, WHITE, True, BODY_FONT)])
    if rows:
        _txt(s, 0.75, 5.05, 9, 1.0, rows, spacing=1.3)
    _txt(s, 0.75, SH - 0.7, 9, 0.4, [[("Informação contábil transformada em decisão.", 11, TEAL, False, BODY_FONT)]])


def _r_kpis(s, sl, num, total):
    _bg(s, WHITE)
    _header(s, sl["title"], sl.get("subtitle") or "", num, total)
    items = sl["data"].get("items") or []
    n = len(items)
    cols = 4 if n > 6 else (3 if n > 4 else max(1, min(n, 3)))
    import math
    rows = math.ceil(n / cols) if n else 1
    L, top, gap = 0.62, 1.95, 0.26
    cw = (12.09 - (cols - 1) * gap) / cols
    ch = min(1.5, (SH - top - 0.7 - (rows - 1) * 0.24) / max(rows, 1))
    for i, it in enumerate(items):
        r, c = divmod(i, cols)
        x = L + c * (cw + gap); y = top + r * (ch + 0.24)
        acc = _ACC.get(it.get("accent", "teal"), TEAL)
        _rect(s, x, y, cw, ch, fill=GRAY_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
        _rect(s, x, y, 0.09, ch, fill=acc)
        _txt(s, x + 0.26, y + 0.16, cw - 0.4, 0.4, [[(it.get("label", ""), 10.5, GRAY, True, BODY_FONT)]])
        _txt(s, x + 0.26, y + 0.52, cw - 0.4, 0.6, [[(str(it.get("value", "")), 21, INK, True, TITLE_FONT)]], anchor=MSO_ANCHOR.TOP)
    _footer(s, "Valores extraídos dos documentos importados. Onde não há dado: 'Dados insuficientes'.")


def _r_barras(s, sl, num, total):
    _bg(s, WHITE)
    _header(s, sl["title"], sl.get("subtitle") or "", num, total)
    d = sl["data"]
    cats = d.get("categories") or []
    series = d.get("series") or []
    if not cats or not series:
        _txt(s, 0.62, 3, 10, 1, [[("Dados insuficientes para o gráfico.", 16, GRAY, False, BODY_FONT)]])
        _footer(s, "")
        return
    cd = CategoryChartData(); cd.categories = cats
    for ser in series:
        cd.add_series(ser.get("name", "Série"), [(_f(v)) for v in ser.get("values", [])])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.62), Inches(2.0),
                            Inches(12.09), Inches(4.4), cd)
    ch = gf.chart; ch.has_legend = len(series) > 1
    if ch.has_legend:
        from pptx.enum.chart import XL_LEGEND_POSITION
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM; ch.legend.include_in_layout = False
    ch.has_title = False
    plot = ch.plots[0]; plot.gap_width = 60
    base_colors = {"teal": TEAL, "indigo": INDIGO, "amber": AMBER, "green": GREEN, "red": RED}
    for si, ser in enumerate(series):
        cser = ch.series[si]
        if d.get("percolor") and ser.get("colors"):
            for pi, cc in enumerate(ser["colors"]):
                try:
                    pt = cser.points[pi]; pt.format.fill.solid()
                    pt.format.fill.fore_color.rgb = base_colors.get(cc, INDIGO)
                except Exception:
                    pass
        else:
            cser.format.fill.solid()
            cser.format.fill.fore_color.rgb = base_colors.get(ser.get("color", "indigo"), INDIGO)
    va = ch.value_axis; va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = GRAY_BD; va.major_gridlines.format.line.width = Pt(0.5)
    va.tick_labels.font.size = Pt(8); va.tick_labels.font.color.rgb = GRAY
    ca = ch.category_axis; ca.tick_labels.font.size = Pt(9); ca.tick_labels.font.bold = True
    ca.tick_labels.font.color.rgb = INK
    plot.has_data_labels = True
    dl = plot.data_labels; dl.show_value = True; dl.show_series_name = False
    dl.show_category_name = False; dl.show_legend_key = False
    dl.number_format = '#,##0' if d.get("unit") == "R$" else '0.00'
    dl.number_format_is_linked = False
    dl.font.size = Pt(9); dl.font.bold = True; dl.font.color.rgb = INK
    _footer(s, "Gráfico nativo editável do PowerPoint.")


def _r_dre(s, sl, num, total):
    _bg(s, WHITE)
    _header(s, sl["title"], sl.get("subtitle") or "", num, total)
    dre = sl["data"].get("dre") or {}
    rows = [
        ("Receita Operacional Bruta", "receita_operacional_bruta", 0, False),
        ("(–) Deduções", "deducoes", 1, False),
        ("(=) Receita Operacional Líquida", "receita_operacional_liquida", 0, True),
        ("(–) Custos", "custos", 1, False),
        ("(=) Lucro Bruto", "lucro_bruto", 0, True),
        ("(–) Despesas Operacionais", "despesas_operacionais", 1, False),
        ("(+) Receitas Financeiras", "receitas_financeiras", 1, False),
        ("(=) Resultado Operacional", "resultado_operacional", 0, True),
        ("(=) Resultado antes dos Tributos", "resultado_antes_tributos", 0, False),
        ("(=) Resultado Líquido", "resultado_liquido", 0, True),
    ]
    from .indicators import num as _n
    data = [(lbl, _n(dre.get(k)), hl) for (lbl, k, _ind, hl) in rows if dre.get(k) is not None]
    if not data:
        _txt(s, 0.62, 3, 10, 1, [[("Dados de DRE insuficientes.", 16, GRAY, False, BODY_FONT)]]); _footer(s, ""); return
    tbl = s.shapes.add_table(len(data) + 1, 2, Inches(0.62), Inches(1.95),
                             Inches(12.09), Inches(0.5 + 0.46 * len(data))).table
    tbl.columns[0].width = Inches(8.8); tbl.columns[1].width = Inches(3.29)
    _cell(tbl.cell(0, 0), "CONTA", INDIGO, WHITE, True, PP_ALIGN.LEFT)
    _cell(tbl.cell(0, 1), "VALOR (R$)", INDIGO, WHITE, True, PP_ALIGN.RIGHT)
    for i, (lbl, val, hl) in enumerate(data, start=1):
        bg = TEAL if hl else (GRAY_LT if i % 2 else WHITE)
        fg = WHITE if hl else INK
        _cell(tbl.cell(i, 0), lbl, bg, fg, hl, PP_ALIGN.LEFT)
        _cell(tbl.cell(i, 1), brl(val), bg, (WHITE if hl else (RED if val < 0 else INK)), True, PP_ALIGN.RIGHT)
    _footer(s, "DRE reconstruída no padrão FELCONT a partir do documento importado (tabela editável).")


def _r_balanco(s, sl, num, total):
    _bg(s, WHITE)
    _header(s, sl["title"], sl.get("subtitle") or "", num, total)
    bal = sl["data"].get("balanco") or {}
    from .indicators import num as _n

    def grp(section, title):
        sec = bal.get(section) or {}
        out = [(title, None, True)]
        for c in (sec.get("contas") or [])[:6]:
            out.append((c.get("descricao", ""), _n(c.get("valor")), False))
        if sec.get("total") is not None:
            out.append((f"Total {title.split()[0].title()}", _n(sec.get("total")), "sub"))
        return out

    left = grp("ativo_circulante", "ATIVO CIRCULANTE") + grp("ativo_nao_circulante", "ATIVO NÃO CIRCULANTE")
    right = grp("passivo_circulante", "PASSIVO CIRCULANTE") + grp("passivo_nao_circulante", "PASSIVO NÃO CIRCULANTE") + grp("patrimonio_liquido", "PATRIMÔNIO LÍQUIDO")

    def render(x, w, title, rows, accent):
        _rect(s, x, 1.95, w, 0.42, fill=accent, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
        _txt(s, x + 0.2, 1.95, w - 0.4, 0.42, [[(title, 12, WHITE, True, BODY_FONT)]], anchor=MSO_ANCHOR.MIDDLE)
        y = 2.5
        for (lbl, val, kind) in rows:
            if kind is True:
                _txt(s, x + 0.1, y, w - 0.2, 0.3, [[(lbl, 10, accent, True, BODY_FONT)]]); y += 0.32
            else:
                bold = kind == "sub"
                _txt(s, x + 0.25, y, w - 2.0, 0.3, [[(lbl, 9.5, INK, bold, BODY_FONT)]])
                _txt(s, x + w - 1.9, y, 1.75, 0.3,
                     [[(brl(val), 9.5, (RED if (val or 0) < 0 else INK), bold, BODY_FONT)]], align=PP_ALIGN.RIGHT)
                y += 0.30
        return y

    render(0.62, 5.9, "ATIVO (Aplicações)", left, INDIGO)
    render(6.83, 5.9, "PASSIVO + PL (Origens)", right, TEAL)
    _footer(s, "Balanço reorganizado no padrão FELCONT — nomenclatura original preservada.")


def _r_lista(s, sl, num, total):
    _bg(s, WHITE)
    _header(s, sl["title"], sl.get("subtitle") or "", num, total)
    items = (sl["data"].get("items") or [])[:6]
    acc = _ACC.get(sl["data"].get("accent", "teal"), TEAL)
    top = 1.95; ch = min(1.35, (SH - top - 0.7) / max(len(items), 1) - 0.18)
    for i, it in enumerate(items):
        y = top + i * (ch + 0.18)
        _rect(s, 0.62, y, 12.09, ch, fill=GRAY_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
        _rect(s, 0.62, y, 0.11, ch, fill=acc)
        tipo = it.get("tipo")
        head = it.get("titulo", "")
        if tipo:
            _txt(s, 0.95, y + 0.14, 11.4, 0.3, [[(tipo.upper() + "  ·  ", 9.5, acc, True, BODY_FONT),
                 (head, 12, INK, True, BODY_FONT)]])
        else:
            _txt(s, 0.95, y + 0.14, 11.4, 0.3, [[(head, 12.5, INK, True, BODY_FONT)]])
        _txt(s, 0.95, y + 0.5, 11.4, ch - 0.5, [[(it.get("texto", ""), 10.5, GRAY, False, BODY_FONT)]], spacing=1.1)
    _footer(s, "Conclusões baseadas exclusivamente nos números importados.")


def _r_conclusao(s, sl, meta):
    _bg(s, INDIGO)
    _rect(s, 0, 0, SW, 0.14, fill=TEAL)
    if os.path.exists(LOGO_PNG):
        s.shapes.add_picture(LOGO_PNG, Inches(0.7), Inches(0.55), width=Inches(2.8))
    d = sl["data"]
    _txt(s, 0.75, 2.0, 9, 0.5, [[("CONCLUSÃO", 14, TEAL, True, BODY_FONT)]])
    _txt(s, 0.75, 2.5, 11.6, 2.6, [[(d.get("resumo") or "Análise concluída.", 18, WHITE, False, TITLE_FONT)]], spacing=1.25)
    flow = ["CONCILIAR", "CORRIGIR", "VALIDAR", "ANALISAR"]
    x = 0.75; y = 5.4; fw = 2.75
    for i, t_ in enumerate(flow):
        _rect(s, x + i * (fw + 0.32), y, fw, 0.9, fill=INDIGO2, line=TEAL, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        _txt(s, x + i * (fw + 0.32), y, fw, 0.9, [[(t_, 15, TEAL, True, TITLE_FONT)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _txt(s, 0.75, SH - 0.6, 11, 0.4, [[("FELCONT — Contabilidade, Finanças e Auditoria", 11, TEAL, True, BODY_FONT)]])


def _cell(cell, text, bg, fg, bold, align):
    cell.fill.solid(); cell.fill.fore_color.rgb = bg
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.12); cell.margin_right = Inches(0.12)
    cell.margin_top = Inches(0.02); cell.margin_bottom = Inches(0.02)
    tf = cell.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(10.5); r.font.bold = bold; r.font.color.rgb = fg; r.font.name = BODY_FONT


def _f(v):
    from .indicators import num
    n = num(v)
    return 0 if n is None else n


_RENDER = {"capa": _r_capa, "kpis": _r_kpis, "resumo": _r_kpis, "barras": _r_barras,
           "dre": _r_dre, "balanco": _r_balanco, "lista": _r_lista, "conclusao": _r_conclusao}


def generate_pptx(slides, meta, out_path):
    prs = Presentation()
    prs.slide_width = Inches(SW); prs.slide_height = Inches(SH)
    blank = prs.slide_layouts[6]
    vis = [sl for sl in slides if sl.get("visible", True)]
    total = len(vis)
    for i, sl in enumerate(vis, start=1):
        s = prs.slides.add_slide(blank)
        fn = _RENDER.get(sl["type"], _r_kpis)
        if sl["type"] in ("capa", "conclusao"):
            fn(s, sl, meta)
        else:
            fn(s, sl, i, total)
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path
