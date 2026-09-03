#!/usr/bin/env python3
"""Gera o .pptx oficial da apresentação CPC 51 | IFRS 18 — Felcont (21 slides, 16:9)."""
import sys
sys.path.insert(0, "/app/frontend/src")
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NAVY = RGBColor(0x07, 0x0A, 0x12)
SURFACE = RGBColor(0x0E, 0x14, 0x24)
CYAN = RGBColor(0x00, 0xE5, 0xFF)
TEAL = RGBColor(0x2D, 0xD4, 0xBF)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
GRAY = RGBColor(0x94, 0xA3, 0xB8)
MUTED = RGBColor(0x64, 0x74, 0x8B)
GREEN = RGBColor(0x10, 0xB9, 0x81)
PURPLE = RGBColor(0x8B, 0x5C, 0xF6)
AMBER = RGBColor(0xF5, 0x9E, 0x0B)
RED = RGBColor(0xF4, 0x3F, 0x5E)
BLUE = RGBColor(0x38, 0xBD, 0xF8)

CAT_HEX = {
    "operacional": CYAN, "investimento": GREEN, "financiamento": PURPLE,
    "impostos": AMBER, "descontinuadas": RED,
}
CAT_NOME = {
    "operacional": "OPERACIONAL", "investimento": "INVESTIMENTO", "financiamento": "FINANCIAMENTO",
    "impostos": "IMPOSTOS S/ RENDA", "descontinuadas": "OP. DESCONTINUADAS",
}

FONT = "Segoe UI"
MONO = "Consolas"
EMBLEM = "/app/frontend/public/brand/emblem.png"
OUT = "/app/frontend/public/downloads/CPC51-IFRS18-Felcont.pptx"

SW, SH = Inches(13.333), Inches(7.5)
prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]

import importlib.util
spec = importlib.util.spec_from_file_location("content", "/dev/null")

# dados (espelho de slidesContent.js)
DRE_ROWS = [
    ("Receita líquida de vendas e serviços", 12500000, "operacional", 0),
    ("Custo dos produtos e serviços (CPV/CSP)", -7800000, "operacional", 0),
    ("Despesas com vendas", -950000, "operacional", 0),
    ("Despesas administrativas", -1400000, "operacional", 0),
    ("Outras receitas operacionais", 180000, "operacional", 0),
    ("LUCRO OPERACIONAL", 2530000, "operacional", 1),
    ("Rendimentos de aplicações financeiras", 220000, "investimento", 0),
    ("Resultado de equivalência patrimonial", 90000, "investimento", 0),
    ("LUCRO ANTES DE FINANCIAMENTO E TRIBUTOS", 2840000, None, 1),
    ("Despesas financeiras — empréstimos", -610000, "financiamento", 0),
    ("Juros sobre arrendamentos (CPC 06 / IFRS 16)", -130000, "financiamento", 0),
    ("LUCRO ANTES DOS TRIBUTOS SOBRE A RENDA", 2100000, None, 1),
    ("IRPJ e CSLL (correntes e diferidos)", -714000, "impostos", 0),
    ("LUCRO DAS OPERAÇÕES CONTINUADAS", 1386000, None, 1),
    ("Resultado de operações descontinuadas (CPC 31)", 74000, "descontinuadas", 0),
    ("LUCRO LÍQUIDO DO EXERCÍCIO", 1460000, None, 2),
]
MAP_ROWS = [
    ("3.01.001", "Receita de prestação de serviços", "operacional", "DRE", "Manter"),
    ("4.01.002", "Custo dos serviços prestados", "operacional", "DRE", "Manter"),
    ("3.02.001", "Rendimentos de aplicações financeiras", "investimento", "DRE", "Reclassificar p/ Investimento"),
    ("3.03.001", "Equivalência patrimonial", "investimento", "DRE", "Reclassificar p/ Investimento"),
    ("4.02.003", "Juros passivos — empréstimos", "financiamento", "DRE", "Reclassificar p/ Financiamento"),
    ("4.02.007", "Juros de arrendamento (CPC 06)", "financiamento", "DRE", "Reclassificar p/ Financiamento"),
    ("4.04.001", "IRPJ / CSLL correntes e diferidos", "impostos", "DRE", "Manter · nova categoria"),
    ("3.05.002", "Resultado de unidade encerrada", "descontinuadas", "DRE", "Avaliar enquadramento CPC 31"),
]
QUESTOES = [
    "O Questor terá parametrização nativa das 5 categorias do CPC 51? Qual o roadmap e o prazo?",
    "A classificação será por conta contábil, por centro de resultado/custo ou por lançamento?",
    "A DRE exibirá automaticamente os subtotais obrigatórios (lucro operacional; antes de financiamento e tributos)?",
    "Como o sistema tratará as MPMs/MPDAs — campos para reconciliação e divulgação em nota?",
    "A DFC (método indireto) partirá do lucro operacional, com juros/dividendos recebidos em investimento e pagos em financiamento?",
    "Haverá rotina para gerar os comparativos de 2026 reexpressos no novo formato em 2027?",
    "É possível exportar a tabela de de-para (conta × categoria CPC 51) para conferência e auditoria?",
]
STEPS = [
    ("01", "Comitê interno CPC 51", "Responsável técnico e sponsor na Felcont."),
    ("02", "Treinar a equipe", "Esta apresentação + texto oficial do CPC 51."),
    ("03", "Mapear clientes impactados", "Carteira por alcance e obrigatoriedade."),
    ("04", "De-para do plano de contas", "Conta × categoria CPC 51, cliente a cliente."),
    ("05", "Validar com o Questor", "Chamado formal + roadmap por escrito."),
    ("06", "Definir MPMs", "Listar medidas e preparar reconciliações."),
    ("07", "Revisar a DFC", "Novo ponto de partida; juros e dividendos."),
    ("08", "Piloto com 2–3 clientes", "Rodar o comparativo de 2026 reexpresso."),
    ("09", "Notas e templates", "Ajustar notas explicativas e modelos."),
    ("10", "Go-live 2027", "Implantação plena + revisão pós-adoção."),
]
RISKS = [
    ("Comparativo 2026 não reexpresso", "Reconstruir um ano de classificação sob pressão de prazo.", "ALTO", RED),
    ("Dependência do fornecedor sem SLA", "Roadmap do Questor sem data vira gargalo fora do nosso controle.", "ALTO", RED),
    ("Classificação incorreta nas categorias", "Subtotais errados e questionamento da auditoria independente.", "ALTO", RED),
    ("Equipe não treinada", "Erros de lançamento replicados em massa na carteira.", "ALTO", RED),
    ("“Outras despesas” genéricas", "Não conformidade com agregação e desagregação.", "MÉDIO", AMBER),
    ("MPMs sem reconciliação", "Medidas divulgadas sem a nota obrigatória de reconciliação.", "MÉDIO", AMBER),
]
NOTES = [
    'Abertura: promessa da reunião — sair sabendo o que muda, quando muda e o que a Felcont fará em 2026. Exemplo: pergunte quem já explicou um EBITDA para cliente — é esse tipo de número que a norma passa a disciplinar.',
    'Contrato da reunião: as 4 dúvidas universais respondidas de antemão. O que é → norma de apresentação, não de mensuração. Quando → 2027 com comparativo 2026 reexpresso. Chamado Questor → sim, para roadmap e de-para. Exemplo: designe ali mesmo o responsável pelo chamado.',
    'IFRS 18 emitida pelo IASB em abril de 2024, substitui a IAS 1 na apresentação. No Brasil: CPC 51, NBC TG 51 (CFC) e Resolução CVM 237 — já é norma brasileira. Exemplo: duas indústrias classificavam juros de arrendamento em lugares diferentes — o "lucro operacional" não era comparável; agora será.',
    'Sobrevoo em três eixos: classificar, apresentar, divulgar. Efeitos em plano de contas, sistemas (Questor), DFC, notas e covenants. Exemplo: covenant atrelado a "lucro operacional" precisa ser relido — o número muda sem a operação mudar.',
    'Regra de ouro: não classificar pelo nome da conta, e sim pela natureza da transação. Operacional é residual. Exceção: bancos/seguradoras/holdings classificam no operacional o que para outros seria investimento/financiamento. Exemplo: rendimento de aplicação de indústria → Investimento; de banco → Operacional.',
    'Subtotal sem ajuste discricionário: tudo que é operacional entra, nada sai "por conveniência". Vitória: comparabilidade. Exemplo: quem excluía depreciação dizendo ser "não recorrente" não poderá mais.',
    'Operacional + investimento = desempenho antes da estrutura de capital e tributos. O "EBIT" padronizado. Exemplo: empresa alavancada vs. sem dívida passam a ser comparáveis na operação.',
    'Percorra a tabela apontando os badges. Valores fictícios. Destaque os dois subtotais obrigatórios e juros de empréstimo/arrendamento após o operacional. Exemplo: aluguel de imóvel não usado pela operação → Investimento.',
    'MPM: subtotal usado em comunicação pública fora das demonstrações, traduzindo a visão da gestão. Não proíbe — disciplina: nota única, reconciliação, efeitos tributários e de PNCL. Exemplo: EBITDA ajustado do release reconciliado até o lucro operacional.',
    'Agregar só o semelhante; desagregar o relevante; "outros" vira residual justificado. Exemplo: qual cliente tem "outras despesas" gorda? Será dos primeiros ajustes.',
    'Natureza, função ou mista — o mais útil. Função ganha dever extra: divulgar depreciação, amortização, benefícios a empregados, impairment e baixas de estoque. Exemplo: quanto de depreciação está no CPV precisará aparecer em nota.',
    'DFC indireta parte do lucro operacional. Fim da discricionariedade: juros/dividendos recebidos → investimento; pagos → financiamento. Exemplo: confirme se o Questor reparametrizará automaticamente.',
    'CPC 51 substitui CPC 26 (R1); NBC TG 51 (nov/2025); Res. CVM 237 revoga 106 e 156; Res. CVM 238 atualiza CPC 03, 06, 15 etc. DVA mantida (Lei 6.404/76). Alcance varia por entidade — validar cliente a cliente.',
    '"2027 parece longe, mas o comparativo de 2027 é 2026." Retrospectiva integral. Exemplo: para publicar o comparativo em março/2027, o de-para precisa rodar desde janeiro/2026.',
    '"Preciso mudar o plano de contas?" Não automaticamente — de-para resolve a maioria dos casos. Pode exigir desdobramentos ("outras despesas"). Depende do roadmap do Questor: chamado primeiro.',
    'Recorte ilustrativo do de-para: conta, categoria, relatório, ação. Exemplo: rendimentos de aplicações mantém o número da conta, mas passa à categoria Investimento.',
    'Perguntas por escrito, com protocolo — resposta documentada vira evidência de planejamento. Exemplo: sem previsão de DFC reparametrizada → item de risco do cliente.',
    'Texto pronto para copiar — troque os colchetes. Um chamado por cliente ou um chamado-mãe da Felcont. Exemplo: abra o chamado-piloto agora e guarde o protocolo.',
    'Governança → diagnóstico → fornecedor → piloto → go-live. Exemplo: piloto com um cliente simples e outro com arrendamentos e aplicações cobre ~90% dos cenários.',
    'O maior risco é o tempo. Glosa de auditoria em março/2027 custa mais que as horas de mapeamento em 2026.',
    'Três frases: a norma existe e tem data; a resposta é método, não pânico; a Felcont sai na frente. Cadeia: Questor → orientação técnica → mapeamento → testes → implementação. Saia com o responsável pelo passo 1 definido.',
]

def fmt(v):
    s = f"{abs(v):,.0f}".replace(",", ".")
    return f"({s})" if v < 0 else s

def slide_bg(s):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY

def box(s, x, y, w, h, fill=SURFACE, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.adjustments[0] = 0.06
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line:
        sh.line.color.rgb = line; sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def tb(s, x, y, w, h):
    sp = s.shapes.add_textbox(x, y, w, h)
    tf = sp.text_frame
    tf.word_wrap = True
    return tf

def para(tf, text, size=12, color=GRAY, bold=False, font=FONT, first=False, align=PP_ALIGN.LEFT, space_after=4):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    r = p.add_run(); r.text = text
    f = r.font; f.size = Pt(size); f.color.rgb = color; f.bold = bold; f.name = font
    return p

def header(s, kicker, title, n, total=21, sub=None):
    para(tb(s, Inches(0.55), Inches(0.32), Inches(9), Inches(0.3)), kicker.upper(), 10, CYAN, True, MONO, True)
    para(tb(s, Inches(0.55), Inches(0.62), Inches(11), Inches(0.8)), title, 28, WHITE, True, FONT, True)
    para(tb(s, Inches(12.2), Inches(0.35), Inches(0.9), Inches(0.3)), f"{n:02d} / {total}", 10, MUTED, False, MONO, True, PP_ALIGN.RIGHT)
    if sub:
        para(tb(s, Inches(0.55), Inches(1.28), Inches(11.5), Inches(0.4)), sub, 12, GRAY, False, FONT, True)

def footer(s):
    s.shapes.add_picture(EMBLEM, Inches(0.55), Inches(7.08), height=Inches(0.26))
    para(tb(s, Inches(0.9), Inches(7.06), Inches(6), Inches(0.3)), "FELCONT · CPC 51 | IFRS 18", 8, MUTED, False, MONO, True)
    para(tb(s, Inches(8.3), Inches(7.06), Inches(4.5), Inches(0.3)), "MATERIAL INTERNO · 2026", 8, MUTED, False, MONO, True, PP_ALIGN.RIGHT)

def new_slide(note_idx):
    s = prs.slides.add_slide(BLANK)
    slide_bg(s)
    s.notes_slide.notes_text_frame.text = NOTES[note_idx]
    return s

def card(s, x, y, w, h, title, body, accent=CYAN, tsize=13, bsize=10.5):
    box(s, x, y, w, h)
    para(tb(s, x + Inches(0.18), y + Inches(0.12), w - Inches(0.36), Inches(0.4)), title, tsize, WHITE, True, FONT, True)
    para(tb(s, x + Inches(0.18), y + Inches(0.55), w - Inches(0.36), h - Inches(0.65)), body, bsize, GRAY, False, FONT, True)

# ---------- 1 · CAPA ----------
s = new_slide(0)
s.shapes.add_picture(EMBLEM, Inches(0.55), Inches(0.42), height=Inches(0.62))
para(tb(s, Inches(1.32), Inches(0.45), Inches(6), Inches(0.4)), "Felcont", 22, WHITE, True, FONT, True)
para(tb(s, Inches(1.33), Inches(0.86), Inches(6), Inches(0.3)), "CONTABILIDADE, FINANÇAS E AUDITORIA", 8, TEAL, False, MONO, True)
para(tb(s, Inches(8.2), Inches(0.5), Inches(4.6), Inches(0.3)), "APRESENTAÇÃO TÉCNICA · 2026", 9, MUTED, False, MONO, True, PP_ALIGN.RIGHT)
para(tb(s, Inches(0.55), Inches(2.0), Inches(11), Inches(0.3)), "NOVA NORMA DE APRESENTAÇÃO DAS DEMONSTRAÇÕES", 11, CYAN, True, MONO, True)
para(tb(s, Inches(0.55), Inches(2.4), Inches(12.2), Inches(1.0)), "CPC 51 | IFRS 18", 54, WHITE, True, FONT, True)
para(tb(s, Inches(0.55), Inches(3.35), Inches(12.2), Inches(1.0)), "O que muda nas demonstrações", 54, CYAN, True, FONT, True)
para(tb(s, Inches(0.55), Inches(4.5), Inches(9.5), Inches(0.7)),
     "O que muda nas demonstrações contábeis e como devemos nos preparar — um guia prático para a equipe Felcont e seus clientes.", 15, GRAY, False, FONT, True)
chips = ["IASB · IFRS 18 · abr/2024", "CFC · NBC TG 51", "CVM · Resolução nº 237", "Vigência · exercícios a partir de 2027"]
cx = 0.55
for c in chips:
    w = Inches(0.16 + 0.093 * len(c))
    box(s, Inches(cx), Inches(5.45), w, Inches(0.34))
    para(tb(s, Inches(cx), Inches(5.5), w, Inches(0.3)), c, 9, GRAY, False, MONO, True, PP_ALIGN.CENTER, 0)
    cx += w.inches + 0.18
footer(s)

# ---------- 2 · QUATRO PERGUNTAS ----------
s = new_slide(1)
header(s, "Resumo executivo", "Quatro perguntas, respostas diretas", 2, sub="As dúvidas que todo cliente e toda equipe fazem — respondidas de antemão.")
qa = [
    ("O que é?", "A IFRS 18 é a nova norma global de apresentação das demonstrações (IASB, abr/2024). O CPC 51 é a versão brasileira — NBC TG 51 (CFC) e Resolução CVM nº 237 — e substitui o CPC 26 (R1). É norma de apresentação: não muda mensuração nem o lucro.", "DETALHES · SLIDE 03"),
    ("O que muda?", "DRE em 5 categorias, dois subtotais obrigatórios, MPMs reconciliadas em nota, fim das “outras despesas” genéricas e DFC reparametrizada. O lucro final não muda — a forma de apresentar, sim.", "SLIDES 04 A 12"),
    ("Quando entra em vigor?", "Exercícios iniciados em ou após 01/01/2027, com aplicação retrospectiva: o comparativo de 2026 já sai no novo formato. Aplicação antecipada permitida. Preparação: 2026.", "LINHA DO TEMPO · SLIDE 14"),
    ("Abrir chamado no Questor para o plano de contas?", "Sim — abra já. Mas não para trocar o plano de contas às cegas: o chamado levanta o roadmap do fornecedor e a parametrização de de-para (conta × categoria), sem quebrar o histórico. [Confirmar com o fornecedor]", "MODELO PRONTO · SLIDE 18"),
]
pos = [(0.55, 1.95), (6.75, 1.95), (0.55, 4.55), (6.75, 4.55)]
for (t, b, ref), (x, y) in zip(qa, pos):
    box(s, Inches(x), Inches(y), Inches(6.05), Inches(2.4))
    para(tb(s, Inches(x + 0.2), Inches(y + 0.14), Inches(5.6), Inches(0.4)), t, 14, WHITE, True, FONT, True)
    para(tb(s, Inches(x + 0.2), Inches(y + 0.58), Inches(5.65), Inches(1.5)), b, 10.5, GRAY, False, FONT, True)
    para(tb(s, Inches(x + 0.2), Inches(y + 2.02), Inches(5.6), Inches(0.3)), ref, 8, MUTED, False, MONO, True)
footer(s)

# ---------- 3 · INTRODUÇÃO ----------
s = new_slide(2)
header(s, "Introdução", "O que é a IFRS 18 — e o que é o CPC 51", 3, sub="A mesma norma, em dois idiomas regulatórios: global (IASB) e brasileiro (CPC / CFC / CVM).")
cols = [
    ("IFRS 18 · IASB", "Emitida em abril de 2024 pelo International Accounting Standards Board. Substitui a IAS 1 na apresentação e divulgação das demonstrações.\n\nObjetivo: comparabilidade — acabar com cada empresa apresentando o resultado de um jeito.", CYAN),
    ("CPC 51 · Brasil", "Pronunciamento do CPC, norma NBC TG 51 do CFC e Resolução CVM nº 237. Substitui o CPC 26 (R1).\n\nConvergência integral com a IFRS 18, com adaptações brasileiras — como a manutenção da DVA.", GREEN),
    ("Por que isso importa", "— Investidores e bancos comparam empresas pelos subtotais da DRE\n— Hoje esses subtotais não são padronizados nem exigidos\n— A norma fixa estrutura, subtotais e divulgações mínimas", BLUE),
]
for i, (t, b, c) in enumerate(cols):
    x = 0.55 + i * 4.2
    box(s, Inches(x), Inches(1.95), Inches(4.0), Inches(3.6))
    para(tb(s, Inches(x + 0.2), Inches(2.12), Inches(3.6), Inches(0.4)), t, 15, c, True, FONT, True)
    para(tb(s, Inches(x + 0.2), Inches(2.6), Inches(3.6), Inches(2.8)), b, 11, GRAY, False, FONT, True)
box(s, Inches(0.55), Inches(5.85), Inches(12.2), Inches(0.7), fill=NAVY, line=CYAN)
para(tb(s, Inches(0.8), Inches(6.02), Inches(11.8), Inches(0.4)),
     "Relação direta: quem aplica CPC 26 (R1) hoje passará a aplicar o CPC 51 — não é uma norma “a mais”, é a substituta.", 12, WHITE, True, FONT, True)
footer(s)

# ---------- 4 · O QUE MUDA ----------
s = new_slide(3)
header(s, "Visão geral", "O que muda na prática", 4, sub="Três eixos — classificar, apresentar e divulgar — com efeitos em toda a cadeia contábil.")
items = [
    ("Classificação", "Receitas e despesas em 5 categorias obrigatórias"),
    ("Subtotais", "Lucro operacional + lucro antes de financiamento e tributos"),
    ("MPMs / MPDAs", "Medidas da administração com nota de reconciliação"),
    ("Agregação", "Fim da linha genérica de “outras despesas”"),
    ("Divulgação", "Notas explicativas mais densas e padronizadas"),
    ("Comparabilidade", "Mesma estrutura para todas as empresas"),
    ("Plano de contas", "De-para conta × categoria (não necessariamente contas novas)"),
    ("Sistemas / ERP", "Parametrização do Questor e rotinas de fechamento"),
    ("DFC", "Ponto de partida e classificação de juros e dividendos"),
    ("Indicadores & covenants", "Contratos atrelados a subtotais precisam ser relidos"),
]
for i, (t, d) in enumerate(items):
    x = 0.55 + (i % 5) * 2.48
    y = 1.95 + (i // 5) * 2.0
    box(s, Inches(x), Inches(y), Inches(2.3), Inches(1.85))
    para(tb(s, Inches(x + 0.14), Inches(y + 0.1), Inches(2.0), Inches(0.3)), f"{i+1:02d}", 10, MUTED, False, MONO, True)
    para(tb(s, Inches(x + 0.14), Inches(y + 0.42), Inches(2.05), Inches(0.55)), t, 11.5, WHITE, True, FONT, True)
    para(tb(s, Inches(x + 0.14), Inches(y + 0.95), Inches(2.05), Inches(0.85)), d, 8.5, GRAY, False, FONT, True)
para(tb(s, Inches(0.55), Inches(6.15), Inches(12), Inches(0.4)),
     "Nenhum item muda o lucro líquido final — muda como ele é decomposto, apresentado e explicado.", 11, AMBER, True, FONT, True)
footer(s)

# ---------- 5 · 5 CATEGORIAS ----------
s = new_slide(4)
header(s, "Nova estrutura da DRE", "Cinco categorias obrigatórias", 5, sub="Toda receita e despesa entra em uma — e somente uma — destas categorias.")
cats = [
    ("01", "Operacional", "Categoria residual: toda receita e despesa das atividades principais que não cair nas demais.", CYAN),
    ("02", "Investimento", "Retornos gerados de forma independente: aplicações, equivalência patrimonial, ativos de investimento.", GREEN),
    ("03", "Financiamento", "Captação de recursos: empréstimos, debêntures e juros de passivos (inclui arrendamentos).", PURPLE),
    ("04", "Impostos s/ renda", "IRPJ e CSLL correntes e diferidos (CPC 32 / IAS 12) e variações cambiais correlatas.", AMBER),
    ("05", "Op. descontinuadas", "Resultado de operações descontinuadas conforme CPC 31 / IFRS 5.", RED),
]
for i, (n, t, d, c) in enumerate(cats):
    x = 0.55 + i * 2.48
    box(s, Inches(x), Inches(1.95), Inches(2.3), Inches(2.9), line=c)
    para(tb(s, Inches(x + 0.16), Inches(2.08), Inches(2.0), Inches(0.4)), n, 18, c, True, MONO, True)
    para(tb(s, Inches(x + 0.16), Inches(2.5), Inches(2.0), Inches(0.6)), t, 12.5, WHITE, True, FONT, True)
    para(tb(s, Inches(x + 0.16), Inches(3.15), Inches(2.0), Inches(1.6)), d, 9, GRAY, False, FONT, True)
box(s, Inches(0.55), Inches(5.1), Inches(6.0), Inches(1.15), fill=NAVY, line=RED)
para(tb(s, Inches(0.75), Inches(5.28), Inches(5.6), Inches(0.9)),
     "Regra de ouro: não se classifica pelo nome da conta, e sim pela natureza da transação.", 11.5, WHITE, True, FONT, True)
box(s, Inches(6.75), Inches(5.1), Inches(6.0), Inches(1.15), fill=NAVY, line=CYAN)
para(tb(s, Inches(6.95), Inches(5.28), Inches(5.6), Inches(0.9)),
     "Exceção: entidades cuja atividade principal é financiar clientes ou investir (bancos, seguradoras) classificam esses resultados no operacional.", 11.5, WHITE, False, FONT, True)
footer(s)

# ---------- 6 · LUCRO OPERACIONAL ----------
s = new_slide(5)
header(s, "Subtotal obrigatório nº 1", "Lucro / prejuízo operacional", 6, sub="A soma de todas as receitas e despesas da categoria operacional — o número que o mercado vai comparar primeiro.")
box(s, Inches(0.55), Inches(2.0), Inches(6.0), Inches(2.0))
para(tb(s, Inches(0.8), Inches(2.15), Inches(5.5), Inches(0.3)), "DEFINIÇÃO", 9, MUTED, True, MONO, True)
para(tb(s, Inches(0.8), Inches(2.5), Inches(5.5), Inches(1.4)),
     "Subtotal que não admite ajuste de critério: tudo que é da operação entra, nada sai “por conveniência”. A depreciação de ativo operacional, por exemplo, fica dentro — sempre.", 13, WHITE, False, FONT, True)
box(s, Inches(0.55), Inches(4.2), Inches(6.0), Inches(2.3), fill=NAVY, line=CYAN)
para(tb(s, Inches(0.8), Inches(4.35), Inches(5.5), Inches(0.3)), "POR QUE MELHORA A COMPARABILIDADE", 9, CYAN, True, MONO, True)
para(tb(s, Inches(0.8), Inches(4.7), Inches(5.5), Inches(1.7)),
     "— Separa o desempenho da operação das decisões de financiamento\n— Comparável entre concorrentes, setores e períodos\n— Passa a ser o ponto de partida da DFC (método indireto)", 12, GRAY, False, FONT, True)
for i, (t, w) in enumerate([("Receitas e despesas da operação", 5.6), ("Sem exclusões “não recorrentes” discricionárias", 4.4), ("Mesma regra para toda empresa e todo período", 3.3)]):
    y = 2.1 + i * 1.15
    box(s, Inches(7.1), Inches(y), Inches(w), Inches(0.85), fill=NAVY, line=CYAN)
    para(tb(s, Inches(7.3), Inches(y + 0.22), Inches(w - 0.3), Inches(0.5)), t, 10.5, GRAY, False, FONT, True)
para(tb(s, Inches(7.1), Inches(5.7), Inches(5.5), Inches(0.4)), "= LUCRO OPERACIONAL", 14, CYAN, True, MONO, True)
footer(s)

# ---------- 7 · LAFT ----------
s = new_slide(6)
header(s, "Subtotal obrigatório nº 2", "Lucro antes de financiamento e tributos", 7, sub="Lucro operacional + categoria de investimento — antes da estrutura de capital e dos tributos.")
box(s, Inches(0.55), Inches(2.1), Inches(5.6), Inches(0.85), line=CYAN)
para(tb(s, Inches(0.8), Inches(2.32), Inches(5.2), Inches(0.4)), "Lucro operacional", 14, WHITE, True, FONT, True)
para(tb(s, Inches(0.55), Inches(3.0), Inches(5.6), Inches(0.4)), "+", 16, MUTED, True, FONT, True, PP_ALIGN.CENTER)
box(s, Inches(0.55), Inches(3.45), Inches(5.6), Inches(0.85), line=GREEN)
para(tb(s, Inches(0.8), Inches(3.67), Inches(5.2), Inches(0.4)), "Categoria de investimento", 14, WHITE, True, FONT, True)
para(tb(s, Inches(0.55), Inches(4.35), Inches(5.6), Inches(0.4)), "=", 16, MUTED, True, FONT, True, PP_ALIGN.CENTER)
box(s, Inches(0.55), Inches(4.8), Inches(5.6), Inches(1.0), fill=NAVY, line=CYAN)
para(tb(s, Inches(0.8), Inches(5.05), Inches(5.2), Inches(0.5)), "Lucro antes de financiamento e tributos", 15, CYAN, True, FONT, True)
card(s, Inches(6.75), Inches(2.1), Inches(6.0), Inches(1.7), "O “EBIT” padronizado",
     "Equivalente padronizado do “EBIT” que analistas calculavam por conta própria — agora exigido, com regra única de composição.", bsize=12)
card(s, Inches(6.75), Inches(4.0), Inches(6.0), Inches(1.7), "Comparabilidade entre estruturas de capital",
     "Empresa alavancada e empresa sem dívida passam a ser comparáveis na operação e no investimento.", bsize=12)
para(tb(s, Inches(6.75), Inches(5.95), Inches(6), Inches(0.4)), "DEPOIS DELE: FINANCIAMENTO → IMPOSTOS S/ RENDA → OP. DESCONTINUADAS", 8.5, MUTED, False, MONO, True)
footer(s)

# ---------- 8 · DRE ----------
s = new_slide(7)
header(s, "Slide-destaque · demonstração", "Exemplo de DRE no novo formato", 8)
tbl_shape = s.shapes.add_table(len(DRE_ROWS) + 1, 3, Inches(0.55), Inches(1.45), Inches(12.25), Inches(5.2))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(7.0); tbl.columns[1].width = Inches(2.6); tbl.columns[2].width = Inches(2.65)
hdr = ["Demonstração do resultado (valores em R$ · FICTÍCIOS)", "Categoria", " "]
for j, h in enumerate(hdr):
    c = tbl.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = NAVY
    tf = c.text_frame; tf.text = h
    p = tf.paragraphs[0]
    if not p.runs:
        continue
    p.runs[0].font.size = Pt(9); p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = MUTED; p.runs[0].font.name = MONO
for i, (label, val, cat, kind) in enumerate(DRE_ROWS, start=1):
    for j in range(3):
        c = tbl.cell(i, j)
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor(0x0B, 0x1D, 0x26) if kind else SURFACE
        c.margin_top = c.margin_bottom = Pt(1)
    c0 = tbl.cell(i, 0).text_frame; c0.text = ("      " if not kind else "") + label
    r0 = c0.paragraphs[0].runs[0]; r0.font.size = Pt(10 if kind else 9.5)
    r0.font.bold = bool(kind); r0.font.name = FONT
    r0.font.color.rgb = WHITE if kind else GRAY
    c1 = tbl.cell(i, 1).text_frame
    if cat:
        c1.text = CAT_NOME[cat]
        r1 = c1.paragraphs[0].runs[0]; r1.font.size = Pt(8); r1.font.bold = True
        r1.font.color.rgb = CAT_HEX[cat]; r1.font.name = MONO
    c2 = tbl.cell(i, 2).text_frame; c2.text = fmt(val)
    c2.paragraphs[0].alignment = PP_ALIGN.RIGHT
    r2 = c2.paragraphs[0].runs[0]; r2.font.size = Pt(10 if kind else 9.5); r2.font.name = MONO
    r2.font.bold = bool(kind); r2.font.color.rgb = CYAN if kind else WHITE
footer(s)

# ---------- 9 · MPMs ----------
s = new_slide(8)
header(s, "Divulgação", "MPMs — medidas definidas pela administração", 9, sub="EBITDA ajustado, lucro recorrente: a norma não proíbe — disciplina.")
box(s, Inches(0.55), Inches(1.95), Inches(6.0), Inches(4.4))
para(tb(s, Inches(0.8), Inches(2.1), Inches(5.5), Inches(0.3)), "O QUE ENQUADRA UMA MEDIDA COMO MPM", 9, MUTED, True, MONO, True)
para(tb(s, Inches(0.8), Inches(2.5), Inches(5.5), Inches(2.2)),
     "1. Subtotal de receitas e despesas;\n2. Usado em comunicação pública FORA das demonstrações (releases, apresentações);\n3. Que transmite a visão da administração sobre o desempenho.", 12, GRAY, False, FONT, True)
box(s, Inches(0.8), Inches(5.35), Inches(5.5), Inches(0.75), fill=NAVY, line=BLUE)
para(tb(s, Inches(1.0), Inches(5.52), Inches(5.2), Inches(0.4)), "Exemplo clássico: “EBITDA ajustado” do release de resultados.", 11, BLUE, True, FONT, True)
regras = [
    ("Nota explicativa única", "Todas as MPMs concentradas em uma única nota."),
    ("Reconciliação obrigatória", "Linha a linha até o subtotal IFRS/CPC mais diretamente comparável."),
    ("Efeitos tributários e de PNCL", "Impostos e participações de não controladores por item."),
    ("Por que a medida é útil", "Explicação da visão da administração e de mudanças de critério."),
]
for i, (t, d) in enumerate(regras):
    y = 1.95 + i * 1.13
    box(s, Inches(6.75), Inches(y), Inches(6.0), Inches(1.0))
    para(tb(s, Inches(6.95), Inches(y + 0.1), Inches(5.6), Inches(0.35)), t, 12, WHITE, True, FONT, True)
    para(tb(s, Inches(6.95), Inches(y + 0.48), Inches(5.6), Inches(0.45)), d, 10, GRAY, False, FONT, True)
footer(s)

# ---------- 10 · AGREGAÇÃO ----------
s = new_slide(9)
header(s, "Transparência", "Agregação e desagregação", 10, sub="Agregar só o que é semelhante; desagregar o que é relevante. “Outros” vira residual — e precisa se justificar.")
box(s, Inches(0.55), Inches(2.3), Inches(4.6), Inches(2.4), fill=NAVY, line=RED)
para(tb(s, Inches(0.8), Inches(2.45), Inches(4.0), Inches(0.3)), "ANTES", 9, RED, True, MONO, True)
para(tb(s, Inches(0.8), Inches(2.85), Inches(4.2), Inches(0.4)), "Outras despesas", 13, GRAY, False, FONT, True)
para(tb(s, Inches(0.8), Inches(3.3), Inches(4.2), Inches(0.6)), "2.800.000", 26, WHITE, True, MONO, True)
para(tb(s, Inches(0.8), Inches(4.05), Inches(4.2), Inches(0.6)), "Uma linha opaca: ninguém sabe o que está dentro.", 10, MUTED, False, FONT, True)
box(s, Inches(5.95), Inches(2.0), Inches(6.8), Inches(4.3), fill=NAVY, line=CYAN)
para(tb(s, Inches(6.2), Inches(2.15), Inches(6.0), Inches(0.3)), "DEPOIS (valores fictícios)", 9, CYAN, True, MONO, True)
depois = [("Manutenção e reparos", 900000), ("Serviços terceirizados", 750000), ("Seguros", 420000), ("Perdas estimadas com créditos (PCLD)", 380000), ("Demais itens não materiais (justificados)", 350000)]
yy = 2.55
for t, v in depois:
    para(tb(s, Inches(6.2), Inches(yy), Inches(4.4), Inches(0.35)), t, 11, GRAY, False, FONT, True)
    para(tb(s, Inches(10.6), Inches(yy), Inches(1.9), Inches(0.35)), fmt(v), 11, WHITE, False, MONO, True, PP_ALIGN.RIGHT)
    yy += 0.52
para(tb(s, Inches(6.2), Inches(yy + 0.1), Inches(4.4), Inches(0.4)), "Total (inalterado)", 12, WHITE, True, FONT, True)
para(tb(s, Inches(10.6), Inches(yy + 0.1), Inches(1.9), Inches(0.4)), "2.800.000", 12, CYAN, True, MONO, True, PP_ALIGN.RIGHT)
para(tb(s, Inches(0.55), Inches(6.5), Inches(12), Inches(0.4)),
     "O valor total não muda — muda a qualidade da informação. Linhas “outras” grandes serão questionadas por auditores e reguladores.", 10.5, MUTED, False, FONT, True)
footer(s)

# ---------- 11 · NATUREZA × FUNÇÃO ----------
s = new_slide(10)
header(s, "Despesas operacionais", "Natureza × função × combinação", 11, sub="O critério mais útil aos usuários — com um dever extra para quem escolher função.")
cols = [
    ("POR NATUREZA", "Tipo do gasto:\n· Matéria-prima e insumos\n· Pessoal e encargos\n· Depreciação e amortização\n· Serviços de terceiros", CYAN),
    ("POR FUNÇÃO", "Área que consome:\n· CPV / custo dos serviços\n· Despesas com vendas\n· Despesas administrativas\n\nDever extra: divulgar em nota depreciação, amortização, benefícios a empregados, impairment e baixas de estoque.", BLUE),
    ("COMBINAÇÃO", "Mista é permitida quando for a mais útil — parte por função, parte por natureza.\n\nO critério precisa ser consistente entre períodos e explicado nas notas.\n\nImpacto direto no plano de contas e no de-para.", GREEN),
]
for i, (t, b, c) in enumerate(cols):
    x = 0.55 + i * 4.2
    box(s, Inches(x), Inches(1.95), Inches(4.0), Inches(4.4))
    para(tb(s, Inches(x + 0.2), Inches(2.12), Inches(3.6), Inches(0.35)), t, 11, c, True, MONO, True)
    para(tb(s, Inches(x + 0.2), Inches(2.55), Inches(3.6), Inches(3.6)), b, 11, GRAY, False, FONT, True)
footer(s)

# ---------- 12 · DFC ----------
s = new_slide(11)
header(s, "Demonstração dos fluxos de caixa", "Impactos na DFC", 12, sub="Novo ponto de partida no método indireto — e fim da liberdade de classificar juros e dividendos.")
box(s, Inches(0.55), Inches(2.1), Inches(5.9), Inches(2.6))
para(tb(s, Inches(0.8), Inches(2.25), Inches(5.4), Inches(0.3)), "PONTO DE PARTIDA (MÉTODO INDIRETO)", 9, MUTED, True, MONO, True)
para(tb(s, Inches(0.8), Inches(2.7), Inches(5.4), Inches(0.5)), "Lucro antes dos tributos  →  LUCRO OPERACIONAL", 13, CYAN, True, FONT, True)
para(tb(s, Inches(0.8), Inches(3.4), Inches(5.4), Inches(1.2)),
     "Menos ajustes artificiais: a reconciliação parte do subtotal que já exclui financiamento e tributos.\n\nEntidades com atividade principal financeira seguem regras próprias.", 11, GRAY, False, FONT, True)
rows = [("Juros recebidos", "investimento"), ("Dividendos recebidos", "investimento"), ("Juros pagos", "financiamento"), ("Dividendos pagos", "financiamento")]
tbl = s.shapes.add_table(5, 2, Inches(6.75), Inches(2.1), Inches(6.0), Inches(2.9)).table
tbl.columns[0].width = Inches(3.3); tbl.columns[1].width = Inches(2.7)
for j, h in enumerate(["Fluxo", "Classificação na DFC"]):
    c = tbl.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = NAVY
    c.text_frame.text = h
    r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = MUTED; r.font.name = MONO
for i, (t, cat) in enumerate(rows, start=1):
    c0 = tbl.cell(i, 0); c0.fill.solid(); c0.fill.fore_color.rgb = SURFACE
    c0.text_frame.text = t
    r0 = c0.text_frame.paragraphs[0].runs[0]; r0.font.size = Pt(11); r0.font.color.rgb = WHITE; r0.font.name = FONT
    c1 = tbl.cell(i, 1); c1.fill.solid(); c1.fill.fore_color.rgb = SURFACE
    c1.text_frame.text = CAT_NOME[cat]
    r1 = c1.text_frame.paragraphs[0].runs[0]; r1.font.size = Pt(9); r1.font.bold = True; r1.font.color.rgb = CAT_HEX[cat]; r1.font.name = MONO
para(tb(s, Inches(6.75), Inches(5.2), Inches(6), Inches(0.7)),
     "Antes: discricionário. Agora: regra única — reparametrização obrigatória no sistema.", 10, MUTED, False, FONT, True)
footer(s)

# ---------- 13 · BRASIL ----------
s = new_slide(12)
header(s, "Cenário regulatório", "Brasil — CPC 51 · NBC TG 51 · CVM 237", 13, sub="Não é “norma de fora”: já é norma brasileira, com cadeia completa de regulação.")
marcos = [
    ("CPC", "CPC 51", "Pronunciamento que substitui o CPC 26 (R1) — Apresentação das Demonstrações Contábeis."),
    ("CFC", "NBC TG 51", "Norma brasileira publicada pelo Conselho Federal de Contabilidade em novembro de 2025."),
    ("CVM", "Resolução nº 237", "Aprova o CPC 51 para companhias abertas e revoga as Resoluções CVM nº 106 e nº 156."),
    ("CVM", "Resolução nº 238", "Torna obrigatório o Documento de Revisão nº 28, atualizando CPC 03, 06, 15 e outros."),
]
for i, (org, doc, d) in enumerate(marcos):
    x = 0.55 + (i % 2) * 6.2
    y = 1.95 + (i // 2) * 1.75
    box(s, Inches(x), Inches(y), Inches(6.0), Inches(1.6))
    para(tb(s, Inches(x + 0.2), Inches(y + 0.12), Inches(1.2), Inches(0.4)), org, 13, CYAN, True, FONT, True)
    para(tb(s, Inches(x + 1.2), Inches(y + 0.15), Inches(4.6), Inches(0.35)), doc, 13, WHITE, True, FONT, True)
    para(tb(s, Inches(x + 1.2), Inches(y + 0.55), Inches(4.6), Inches(0.95)), d, 10, GRAY, False, FONT, True)
box(s, Inches(0.55), Inches(5.65), Inches(12.2), Inches(0.85))
para(tb(s, Inches(0.8), Inches(5.8), Inches(11.8), Inches(0.6)),
     "Adaptações mantidas: DVA (Lei 6.404/76) e particularidades da Lei das S.A.  ·  Alcance da obrigatoriedade varia por tipo de entidade — validar cliente a cliente [Confirmar na fonte].", 11, WHITE, False, FONT, True)
footer(s)

# ---------- 14 · VIGÊNCIA ----------
s = new_slide(13)
header(s, "Vigência e transição", "2026 prepara · 2027 adota", 14, sub="Aplicação antecipada permitida. Transição retrospectiva: o comparativo de 2026 entra no novo formato.")
pts = [
    ("abr/2024", "Emissão da IFRS 18", "IASB publica a norma; início do debate global.", False),
    ("2025", "CPC 51 · NBC TG 51 · CVM 237", "Convergência brasileira formalizada.", False),
    ("2026", "Ano de preparação", "Mapeamento, de-para, pilotos — e o ano que virará comparativo reexpresso.", True),
    ("2027", "Adoção obrigatória", "Exercícios a partir de 01/01/2027. Aplicação retrospectiva integral.", True),
]
for i, (ano, t, d, on) in enumerate(pts):
    x = 0.55 + i * 3.12
    box(s, Inches(x), Inches(2.3), Inches(2.9), Inches(2.6), line=CYAN if on else None)
    para(tb(s, Inches(x + 0.18), Inches(2.45), Inches(2.5), Inches(0.4)), ano, 16, CYAN if on else MUTED, True, MONO, True)
    para(tb(s, Inches(x + 0.18), Inches(2.95), Inches(2.5), Inches(0.6)), t, 12, WHITE, True, FONT, True)
    para(tb(s, Inches(x + 0.18), Inches(3.6), Inches(2.55), Inches(1.2)), d, 9.5, GRAY, False, FONT, True)
box(s, Inches(0.55), Inches(5.35), Inches(12.2), Inches(1.0), fill=NAVY, line=AMBER)
para(tb(s, Inches(0.8), Inches(5.55), Inches(11.8), Inches(0.7)),
     "Por que 2026 importa: o comparativo publicado em 2027 é o exercício de 2026 — quem não reexpressar durante o ano vai reconstruir 12 meses de classificação sob pressão de prazo.", 12, WHITE, False, FONT, True)
footer(s)

# ---------- 15 · QUESTOR ----------
s = new_slide(14)
header(s, "Sistemas · ERP Questor", "“Precisamos mudar o plano de contas?”", 15, sub="Resposta ponderada: não automaticamente. A norma exige classificação e apresentação — não necessariamente contas novas.")
box(s, Inches(0.55), Inches(1.95), Inches(6.0), Inches(3.6), line=GREEN)
para(tb(s, Inches(0.8), Inches(2.1), Inches(5.5), Inches(0.3)), "CAMINHO PROVÁVEL — SEM QUEBRAR HISTÓRICO", 9, GREEN, True, MONO, True)
para(tb(s, Inches(0.8), Inches(2.5), Inches(5.5), Inches(2.9)),
     "— Tabela de de-para: conta existente × categoria CPC 51\n— Parametrização no Questor por conta ou centro de resultado\n— Relatórios DRE/DFC remontados a partir do de-para\n— Histórico e comparabilidade interna preservados", 11.5, GRAY, False, FONT, True)
box(s, Inches(6.75), Inches(1.95), Inches(6.0), Inches(3.6), line=AMBER)
para(tb(s, Inches(7.0), Inches(2.1), Inches(5.5), Inches(0.3)), "QUANDO PODE EXIGIR CONTAS NOVAS", 9, AMBER, True, MONO, True)
para(tb(s, Inches(7.0), Inches(2.5), Inches(5.5), Inches(2.9)),
     "— Contas “guarda-chuva” (ex.: outras despesas genéricas) pedem desdobramento\n— Despesas por função exigem rateio de natureza rastreável\n— MPMs pedem campos/tags para reconciliação automática\n— Decisão final depende do roadmap do fornecedor", 11.5, GRAY, False, FONT, True)
para(tb(s, Inches(0.55), Inches(5.9), Inches(12.2), Inches(0.5)),
     "Conduta Felcont: primeiro o chamado formal ao Questor, depois qualquer alteração de plano de contas — nunca antes. [Confirmar com o fornecedor]", 12, WHITE, True, FONT, True)
footer(s)

# ---------- 16 · MAPEAMENTO ----------
s = new_slide(15)
header(s, "De-para · exemplo fictício", "Tabela de mapeamento de contas", 16)
tbl = s.shapes.add_table(len(MAP_ROWS) + 1, 5, Inches(0.55), Inches(1.5), Inches(12.25), Inches(4.6)).table
widths = [1.4, 4.0, 2.5, 1.15, 3.2]
for j, wd in enumerate(widths):
    tbl.columns[j].width = Inches(wd)
for j, h in enumerate(["Conta", "Descrição", "Categoria CPC 51", "Relatório", "Ação"]):
    c = tbl.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = NAVY
    c.text_frame.text = h
    r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(9); r.font.bold = True; r.font.color.rgb = MUTED; r.font.name = MONO
for i, (conta, desc, cat, rel, acao) in enumerate(MAP_ROWS, start=1):
    vals = [conta, desc, CAT_NOME[cat], rel, acao]
    colors = [CYAN, WHITE, CAT_HEX[cat], MUTED, GRAY]
    fonts = [MONO, FONT, MONO, MONO, FONT]
    for j, v in enumerate(vals):
        c = tbl.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = SURFACE
        c.text_frame.text = v
        r = c.text_frame.paragraphs[0].runs[0]
        r.font.size = Pt(9.5); r.font.color.rgb = colors[j]; r.font.name = fonts[j]
        r.font.bold = j == 2
para(tb(s, Inches(0.55), Inches(6.35), Inches(12), Inches(0.4)),
     "Recorte ilustrativo — o de-para real nasce do plano de contas de cada cliente, validado com a parametrização do Questor. [Confirmar com o fornecedor]", 10, MUTED, False, FONT, True)
footer(s)

# ---------- 17 · PERGUNTAS QUESTOR ----------
s = new_slide(16)
header(s, "Fornecedor · preparação", "Perguntas a fazer ao suporte Questor", 17, sub="Perguntas objetivas, por escrito — a resposta documentada vira evidência de planejamento.")
for i, q in enumerate(QUESTOES):
    y = 1.85 + i * 0.66
    box(s, Inches(0.55), Inches(y), Inches(12.2), Inches(0.56))
    para(tb(s, Inches(0.75), Inches(y + 0.08), Inches(0.5), Inches(0.4)), f"{i+1:02d}", 11, CYAN, True, MONO, True)
    para(tb(s, Inches(1.3), Inches(y + 0.1), Inches(11.2), Inches(0.42)), q, 10, GRAY, False, FONT, True)
para(tb(s, Inches(0.55), Inches(6.55), Inches(12), Inches(0.4)),
     "Regra da casa: nada de telefone — chamado formal, com protocolo e resposta por escrito.", 10.5, WHITE, True, FONT, True)
footer(s)

# ---------- 18 · MODELO CHAMADO ----------
s = new_slide(17)
header(s, "Pronto para usar", "Modelo de chamado ao Questor", 18, sub="Troque os colchetes pelos dados do cliente e envie. Um chamado por cliente — ou um chamado-mãe da Felcont.")
TICKET = """Assunto: CPC 51 / IFRS 18 — Roadmap de adequação da DRE, DFC e parametrização de categorias

Prezada equipe de suporte Questor,

Somos responsáveis pela contabilidade da empresa [RAZÃO SOCIAL], CNPJ [00.000.000/0000-00], e estamos planejando a adequação ao CPC 51 / NBC TG 51 (IFRS 18), obrigatório para exercícios iniciados a partir de 01/01/2027, com comparativos de 2026 reexpressos.

Solicitamos, por escrito: 1. Roadmap e prazo para parametrização das 5 categorias da DRE; 2. Forma de classificação (conta, centro de resultado ou lançamento); 3. Geração automática dos subtotais obrigatórios; 4. Tratamento das MPMs/MPDAs e reconciliações em nota; 5. Adequação da DFC (método indireto a partir do lucro operacional; juros/dividendos recebidos em investimento; pagos em financiamento); 6. Rotina para comparativos de 2026 reexpressos; 7. Exportação da tabela de de-para (conta × categoria) para auditoria.

Ficamos à disposição para reunião técnica e solicitamos o número de protocolo deste chamado.

Atenciosamente,
[NOME] — Felcont Contabilidade, Finanças e Auditoria · [E-MAIL] · [TELEFONE]"""
box(s, Inches(0.55), Inches(1.9), Inches(12.2), Inches(4.9), fill=RGBColor(0x0B, 0x10, 0x1E))
para(tb(s, Inches(0.9), Inches(2.1), Inches(11.6), Inches(4.5)), TICKET, 10, GRAY, False, FONT, True)
footer(s)

# ---------- 19 · PLANO 10 PASSOS ----------
s = new_slide(18)
header(s, "Execução", "Plano de ação Felcont 2026 — 10 passos", 19, sub="Governança → diagnóstico → fornecedor → piloto → go-live. Cada passo com dono e prazo.")
for i, (n, t, d) in enumerate(STEPS):
    x = 0.55 + (i % 5) * 2.48
    y = 1.95 + (i // 5) * 2.05
    box(s, Inches(x), Inches(y), Inches(2.3), Inches(1.9))
    para(tb(s, Inches(x + 0.14), Inches(y + 0.1), Inches(2.0), Inches(0.35)), n, 14, CYAN, True, MONO, True)
    para(tb(s, Inches(x + 0.14), Inches(y + 0.5), Inches(2.05), Inches(0.6)), t, 11, WHITE, True, FONT, True)
    para(tb(s, Inches(x + 0.14), Inches(y + 1.1), Inches(2.05), Inches(0.75)), d, 8.5, GRAY, False, FONT, True)
para(tb(s, Inches(0.55), Inches(6.25), Inches(12), Inches(0.4)),
     "Piloto recomendado: um cliente simples + um com arrendamentos e aplicações — cobre ~90% dos cenários da carteira.", 10.5, MUTED, False, FONT, True)
footer(s)

# ---------- 20 · RISCOS ----------
s = new_slide(19)
header(s, "Gestão de riscos", "O custo de deixar para a última hora", 20, sub="O maior risco não é a norma ser difícil — é o tempo.")
for i, (t, d, nivel, c) in enumerate(RISKS):
    x = 0.55 + (i % 3) * 4.2
    y = 1.95 + (i // 3) * 1.95
    box(s, Inches(x), Inches(y), Inches(4.0), Inches(1.8), line=c)
    para(tb(s, Inches(x + 0.16), Inches(y + 0.1), Inches(2.9), Inches(0.55)), t, 11.5, WHITE, True, FONT, True)
    para(tb(s, Inches(x + 3.1), Inches(y + 0.14), Inches(0.8), Inches(0.3)), nivel, 9, c, True, MONO, True, PP_ALIGN.RIGHT)
    para(tb(s, Inches(x + 0.16), Inches(y + 0.75), Inches(3.7), Inches(0.95)), d, 9.5, GRAY, False, FONT, True)
para(tb(s, Inches(0.55), Inches(6.1), Inches(12), Inches(0.5)),
     "Uma glosa de auditoria em março de 2027 custa mais caro — em retrabalho e prazo regulatório — do que as horas de mapeamento em 2026.", 11, MUTED, False, FONT, True)
footer(s)

# ---------- 21 · CONCLUSÃO ----------
s = new_slide(20)
header(s, "Fechamento", "Conclusão e próximos passos", 21, sub="Transformar exigência regulatória em diferencial consultivo para os clientes.")
final = [
    ("01", "A norma já existe e tem data", "2027 obrigatória, com comparativo de 2026 reexpresso — retrospectiva integral."),
    ("02", "A resposta é método, não pânico", "De-para, validação com o Questor e piloto — não plano de contas novo às cegas."),
    ("03", "A Felcont sai na frente", "Quem orienta o cliente em 2026 vende tranquilidade — não retrabalho em 2027."),
]
for i, (n, t, d) in enumerate(final):
    y = 1.95 + i * 1.35
    box(s, Inches(0.55), Inches(y), Inches(6.6), Inches(1.2))
    para(tb(s, Inches(0.75), Inches(y + 0.12), Inches(0.6), Inches(0.4)), n, 13, CYAN, True, MONO, True)
    para(tb(s, Inches(1.35), Inches(y + 0.1), Inches(5.6), Inches(0.4)), t, 13, WHITE, True, FONT, True)
    para(tb(s, Inches(1.35), Inches(y + 0.55), Inches(5.6), Inches(0.55)), d, 10.5, GRAY, False, FONT, True)
box(s, Inches(7.5), Inches(1.95), Inches(5.25), Inches(4.3), fill=NAVY, line=CYAN)
para(tb(s, Inches(7.7), Inches(2.1), Inches(4.8), Inches(0.3)), "CADEIA DE EXECUÇÃO", 9, CYAN, True, MONO, True)
cadeia = ["Questor", "Orientação técnica", "Mapeamento", "Testes", "Implementação"]
for i, p in enumerate(cadeia):
    y = 2.5 + i * 0.68
    box(s, Inches(7.7), Inches(y), Inches(4.85), Inches(0.55))
    para(tb(s, Inches(7.9), Inches(y + 0.1), Inches(4.4), Inches(0.35)), f"{i+1:02d}  ·  {p}", 11, WHITE, True, FONT, True)
para(tb(s, Inches(7.7), Inches(6.0), Inches(4.8), Inches(0.3)), "FELCONT · CONTABILIDADE, FINANÇAS E AUDITORIA", 8, TEAL, False, MONO, True)
footer(s)

prs.save(OUT)
print("OK", OUT, "slides:", len(prs.slides.__iter__.__self__._sldIdLst))
