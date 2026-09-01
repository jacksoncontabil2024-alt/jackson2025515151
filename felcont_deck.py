# -*- coding: utf-8 -*-
"""Gerador da Apresentacao Gerencial FELCONT (.pptx) - 16:9, pt-BR, 9 slides."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData
from pptx.oxml.ns import qn

# ------------------------------------------------------------------ paleta
BLUE      = RGBColor(0x04, 0xB7, 0xAF)   # turquesa primario FELCONT
BLUE_DK   = RGBColor(0x06, 0x8B, 0x84)
NAVY      = RGBColor(0x32, 0x2F, 0x6A)   # indigo da marca (fundo escuro)
NAVY2     = RGBColor(0x3E, 0x3A, 0x82)
ORANGE    = RGBColor(0xE8, 0xA1, 0x3A)   # ambar (reducoes / atencao)
GREEN     = RGBColor(0x57, 0xB1, 0x4A)
RED       = RGBColor(0xD6, 0x45, 0x3F)
RED_BG    = RGBColor(0xFB, 0xEC, 0xEC)
GRAY_TX   = RGBColor(0x26, 0x24, 0x45)
GRAY      = RGBColor(0x6B, 0x6E, 0x86)
GRAY_LT   = RGBColor(0xF2, 0xF3, 0xF8)
GRAY_BD   = RGBColor(0xDE, 0xE0, 0xEC)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
AMBER_BG  = RGBColor(0xFF, 0xF6, 0xE6)
LOGO_PNG  = "/app/assets_in/a2.png"

TITLE_FONT = "Roboto"
BODY_FONT  = "Open Sans"

EMU_IN = 914400
prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = 13.333, 7.5

# ------------------------------------------------------------------ helpers
def brl(v, cents=True):
    n = f"{abs(v):,.2f}" if cents else f"{abs(v):,.0f}"
    n = n.replace(",", "§").replace(".", ",").replace("§", ".")
    sign = "-" if v < 0 else ""
    return f"{sign}R$ {n}"

def slide(bg=WHITE):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid(); r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s

def rect(s, l, t, w, h, fill=None, line=None, lw=1.0, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = s.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(lw)
    sp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0] = radius
        except Exception: pass
    return sp

def txt(s, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
        spacing=1.0, sp_after=2):
    """runs: list of (text,size,color,bold,font) or list of such lists (paragraphs)."""
    box = s.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{m}", 0)
    paras = runs if runs and isinstance(runs[0], list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = spacing; p.space_after = Pt(sp_after)
        p.space_before = Pt(0)
        for (text, size, color, bold, font) in para:
            r = p.add_run(); r.text = text
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = font
    return box

def logo(s, l, t, dark=False):
    sub = WHITE if dark else GRAY
    txt(s, l, t, 5.5, 0.9, [
        [("FEL", 26, BLUE, True, TITLE_FONT), ("CONT", 26, (WHITE if dark else GRAY_TX), True, TITLE_FONT)],
        [("CONTABILIDADE  ·  FINANÇAS  ·  AUDITORIA", 9, sub, False, BODY_FONT)],
    ], spacing=1.0, sp_after=0)

def header(s, kicker, title, num):
    rect(s, 0, 0, SW, 0.14, fill=BLUE)
    txt(s, 0.62, 0.5, 10, 0.4, [[(kicker.upper(), 12, BLUE, True, BODY_FONT)]])
    txt(s, 0.62, 0.86, 11.5, 0.9, [[(title, 28, GRAY_TX, True, TITLE_FONT)]])
    rect(s, 0.63, 1.62, 1.3, 0.06, fill=ORANGE)
    txt(s, SW-1.6, 0.55, 1.0, 0.5, [[(f"{num:02d}", 22, GRAY_BD, True, TITLE_FONT)],
        [("/ 09", 10, GRAY_BD, False, BODY_FONT)]], align=PP_ALIGN.RIGHT)

def footer(s, note):
    rect(s, 0.62, SH-0.52, 12.09, 0.012, fill=GRAY_BD)
    txt(s, 0.62, SH-0.44, 11.0, 0.35, [[(note, 8.5, GRAY, False, BODY_FONT)]])
    txt(s, SW-2.4, SH-0.44, 1.78, 0.35, [[("FELCONT", 8.5, BLUE, True, BODY_FONT)]], align=PP_ALIGN.RIGHT)

def _hide_label(ser, idx):
    dLbls = ser.find(qn('c:dLbls'))
    if dLbls is None:
        dLbls = ser.makeelement(qn('c:dLbls'), {})
        cat = ser.find(qn('c:cat'))
        ser.insert(list(ser).index(cat), dLbls)
    dLbl = ser.makeelement(qn('c:dLbl'), {})
    dLbl.append(ser.makeelement(qn('c:idx'), {'val': str(idx)}))
    dLbl.append(ser.makeelement(qn('c:delete'), {'val': '1'}))
    dLbls.insert(0, dLbl)

# ================================================================== SLIDE 1
s = slide(NAVY)
# faixa lateral decorativa
rect(s, 8.7, 0, 4.633, SH, fill=NAVY2)
rect(s, 8.7, 0, 0.10, SH, fill=BLUE)
# barras ascendentes (elemento financeiro sutil)
bar_x = [9.15, 9.85, 10.55, 11.25, 11.95]
bar_h = [1.1, 1.7, 1.45, 2.35, 3.0]
bar_c = [BLUE_DK, BLUE, BLUE_DK, BLUE, BLUE]
for x, h, c in zip(bar_x, bar_h, bar_c):
    rect(s, x, 5.7-h, 0.5, h, fill=c, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.18)
# linha de tendencia
pts = [(bar_x[i]+0.25, 5.7-bar_h[i]-0.15) for i in range(5)]
for i in range(4):
    ln = s.shapes.add_connector(2, Inches(pts[i][0]), Inches(pts[i][1]), Inches(pts[i+1][0]), Inches(pts[i+1][1]))
    ln.line.color.rgb = ORANGE; ln.line.width = Pt(2.25); ln.shadow.inherit = False
for (px, py) in pts:
    d = rect(s, px-0.06, py-0.06, 0.12, 0.12, fill=WHITE, shape=MSO_SHAPE.OVAL)

s.shapes.add_picture(LOGO_PNG, Inches(0.7), Inches(0.6), width=Inches(3.0))
txt(s, 0.75, 2.55, 8.0, 0.5, [[("ANÁLISE GERENCIAL", 16, BLUE, True, BODY_FONT)]])
txt(s, 0.75, 3.05, 8.0, 1.7, [
    [("Demonstração do", 40, WHITE, True, TITLE_FONT)],
    [("Resultado do Exercício", 40, WHITE, True, TITLE_FONT)],
], spacing=1.05)
rect(s, 0.78, 4.85, 1.5, 0.06, fill=ORANGE)
txt(s, 0.75, 5.15, 8.0, 1.2, [
    [("Cliente:  ", 14, GRAY, False, BODY_FONT), ("Ronaldo Donadon", 14, WHITE, True, BODY_FONT)],
    [("Período:  ", 14, GRAY, False, BODY_FONT), ("Janeiro a Julho de 2026", 14, WHITE, True, BODY_FONT)],
], spacing=1.3)
txt(s, 0.75, SH-0.7, 9, 0.4, [[("Informação contábil transformada em decisão.", 11, BLUE, False, BODY_FONT)]])

# ================================================================== SLIDE 2
s = slide(WHITE)
header(s, "Panorama", "Visão Geral do Resultado", 2)
cards = [
    ("Receita Op. Bruta", 366795.54, BLUE, "Faturamento bruto do período"),
    ("(–) Deduções", 27061.34, ORANGE, "Impostos e deduções s/ receita"),
    ("Receita Líquida", 339734.20, BLUE, "Base operacional líquida"),
    ("(–) Custos", 301937.44, ORANGE, "Custos dos serviços/produtos"),
    ("Lucro Bruto", 37796.76, GREEN, "Receita líquida – custos"),
    ("(–) Despesas Operacionais", 394936.02, ORANGE, "Vendas, adm., financ. e tributárias"),
]
L, top, gap = 0.62, 1.95, 0.28
cw = (12.09 - 2*gap) / 3
ch = 1.28
for i, (t_, v, acc, d) in enumerate(cards):
    r, c = divmod(i, 3)
    x = L + c*(cw+gap); y = top + r*(ch+0.24)
    rect(s, x, y, cw, ch, fill=GRAY_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    rect(s, x, y, 0.09, ch, fill=acc)
    txt(s, x+0.28, y+0.16, cw-0.4, 0.35, [[(t_, 10.5, GRAY, True, BODY_FONT)]])
    txt(s, x+0.28, y+0.5, cw-0.4, 0.45, [[(brl(v), 22, GRAY_TX, True, TITLE_FONT)]])
    txt(s, x+0.28, y+0.98, cw-0.4, 0.28, [[(d, 8.5, GRAY, False, BODY_FONT)]])
# card destaque prejuizo
y = top + 2*(ch+0.24) + 0.02
rect(s, L, y, 12.09, 1.32, fill=RED_BG, line=RED, lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
rect(s, L, y, 0.14, 1.32, fill=RED)
rect(s, L+0.4, y+0.42, 0.5, 0.5, fill=RED, shape=MSO_SHAPE.OVAL)
txt(s, L+0.4, y+0.46, 0.5, 0.42, [[("!", 22, WHITE, True, TITLE_FONT)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
txt(s, L+1.15, y+0.24, 7.0, 0.9, [
    [("RESULTADO LÍQUIDO DO PERÍODO", 12, RED, True, BODY_FONT)],
    [("Prejuízo apurado — exige atenção gerencial", 10.5, GRAY_TX, False, BODY_FONT)],
], spacing=1.15, anchor=MSO_ANCHOR.MIDDLE)
txt(s, 7.6, y+0.18, 4.9, 1.0, [
    [("PREJUÍZO", 12, RED, True, BODY_FONT)],
    [(brl(357139.26).replace("R$", "–R$"), 30, RED, True, TITLE_FONT)],
], align=PP_ALIGN.RIGHT, spacing=1.05, anchor=MSO_ANCHOR.MIDDLE)
footer(s, "Despesas Operacionais incluem Vendas, Administrativas, Financeiras (líquidas) e Tributárias. Fonte: DRE e Balancete do período.")

# ================================================================== SLIDE 3
s = slide(WHITE)
header(s, "Ponte do Resultado", "Formação do Resultado", 3)
cd = CategoryChartData()
cd.categories = ["Receita\nBruta", "(–) Deduções", "Receita\nLíquida", "(–) Custos", "Lucro\nBruto", "(–) Despesas", "Prejuízo"]
flow = [
    (0, 366795.54), (339734.20, 366795.54), (0, 339734.20), (37796.76, 339734.20),
    (0, 37796.76), (-357139.26, 37796.76), (-357139.26, 0),
]
base, pos, neg = [], [], []
for b, t in flow:
    base.append(b if b > 0 else 0)
    pos.append((t - b) if b > 0 else (t if t > 0 else 0))
    neg.append(b if b < 0 else 0)
cd.add_series("base", base)
cd.add_series("pos", pos)
cd.add_series("neg", neg)
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(0.62), Inches(2.0),
                        Inches(12.09), Inches(4.1), cd)
chart = gf.chart
chart.has_legend = False
chart.has_title = False
plot = chart.plots[0]
plot.gap_width = 55
s_base, s_pos, s_neg = chart.series[0], chart.series[1], chart.series[2]
s_base.format.fill.background(); s_base.format.line.fill.background()
pos_colors = [BLUE, ORANGE, BLUE, ORANGE, GREEN, ORANGE, GRAY_LT]
for i, col in enumerate(pos_colors):
    p = s_pos.points[i]; p.format.fill.solid(); p.format.fill.fore_color.rgb = col
    p.format.line.color.rgb = WHITE; p.format.line.width = Pt(0.75)
for i in range(7):
    p = s_neg.points[i]
    p.format.fill.solid(); p.format.fill.fore_color.rgb = (ORANGE if i == 5 else RED)
    p.format.line.color.rgb = WHITE; p.format.line.width = Pt(0.75)
# eixos
va = chart.value_axis
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = GRAY_BD
va.major_gridlines.format.line.width = Pt(0.5)
va.tick_labels.font.size = Pt(8); va.tick_labels.font.name = BODY_FONT
va.tick_labels.font.color.rgb = GRAY
va.tick_labels.number_format = '#,##0'; va.tick_labels.number_format_is_linked = False
va.format.line.fill.background()
ca = chart.category_axis
ca.tick_labels.font.size = Pt(9); ca.tick_labels.font.bold = True
ca.tick_labels.font.name = BODY_FONT; ca.tick_labels.font.color.rgb = GRAY_TX
ca.format.line.color.rgb = GRAY_BD
# rotulos de dados (texto explicito -> sem nome de serie/categoria)
plot.has_data_labels = True
dl = plot.data_labels
dl.show_value = False; dl.show_series_name = False; dl.show_category_name = False
dl.show_legend_key = False; dl.show_percentage = False

def set_label(ser_obj, idx, text, color=GRAY_TX):
    d = ser_obj.points[idx].data_label
    tf = d.text_frame; tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size = Pt(9); r.font.bold = True
            r.font.name = BODY_FONT; r.font.color.rgb = color

labels_pos = {0: (brl(366795.54), GRAY_TX), 1: (brl(27061.34), GRAY_TX),
              2: (brl(339734.20), GRAY_TX), 3: (brl(301937.44), WHITE),
              4: (brl(37796.76), GRAY_TX)}
for i, (t_, c_) in labels_pos.items():
    set_label(s_pos, i, t_, c_)
set_label(s_neg, 5, brl(394936.02), WHITE)
set_label(s_neg, 6, brl(-357139.26), WHITE)
footer(s, "Cascata: Receita Bruta → Deduções → Receita Líquida → Custos → Lucro Bruto → (–) Despesas Op. → Prejuízo.")

# ================================================================== SLIDE 4
s = slide(WHITE)
header(s, "Comparativo", "Receita x Custos e Despesas", 4)
cd = CategoryChartData()
cd.categories = ["Receita Líquida", "Custos", "Despesas Operacionais"]
cd.add_series("Valores", (339734.20, 301937.44, 394936.02))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.62), Inches(2.0),
                        Inches(7.3), Inches(4.1), cd)
chart = gf.chart
chart.has_legend = False; chart.has_title = False
plot = chart.plots[0]; plot.gap_width = 70
ser = chart.series[0]
for i, col in enumerate([BLUE, ORANGE, RED]):
    p = ser.points[i]; p.format.fill.solid(); p.format.fill.fore_color.rgb = col
plot.has_data_labels = True
dl = plot.data_labels
dl.show_value = True; dl.show_series_name = False; dl.show_category_name = False
dl.show_legend_key = False; dl.show_percentage = False
dl.number_format = '"R$ "#,##0.00'; dl.number_format_is_linked = False
dl.font.size = Pt(10); dl.font.bold = True; dl.font.name = BODY_FONT; dl.font.color.rgb = GRAY_TX
va = chart.value_axis
va.has_major_gridlines = True
va.major_gridlines.format.line.color.rgb = GRAY_BD
va.major_gridlines.format.line.width = Pt(0.5)
va.tick_labels.font.size = Pt(8); va.tick_labels.number_format = '#,##0'
va.tick_labels.number_format_is_linked = False
va.tick_labels.font.color.rgb = GRAY; va.format.line.fill.background()
ca = chart.category_axis
ca.tick_labels.font.size = Pt(9.5); ca.tick_labels.font.bold = True
ca.tick_labels.font.color.rgb = GRAY_TX; ca.format.line.color.rgb = GRAY_BD
# conclusao
bx = 8.2
rect(s, bx, 2.0, 4.5, 4.1, fill=GRAY_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
rect(s, bx, 2.0, 4.5, 0.09, fill=RED)
txt(s, bx+0.35, 2.35, 3.8, 0.4, [[("LEITURA DO GRÁFICO", 11, RED, True, BODY_FONT)]])
txt(s, bx+0.35, 2.85, 3.85, 2.7, [
    [("A soma de Custos e Despesas ", 12.5, GRAY_TX, False, BODY_FONT),
     ("supera em muito ", 12.5, RED, True, BODY_FONT),
     ("a Receita Líquida do período.", 12.5, GRAY_TX, False, BODY_FONT)],
    [("", 6, GRAY_TX, False, BODY_FONT)],
    [("Custos + Despesas: ", 11, GRAY, True, BODY_FONT),
     (brl(696873.46), 11, GRAY_TX, True, BODY_FONT)],
    [("Receita Líquida: ", 11, GRAY, True, BODY_FONT),
     (brl(339734.20), 11, GRAY_TX, True, BODY_FONT)],
    [("", 6, GRAY_TX, False, BODY_FONT)],
    [("Estrutura de gastos incompatível com a geração de receita atual — resultado operacional deficitário.",
      11, GRAY_TX, False, BODY_FONT)],
], spacing=1.2, sp_after=3)
footer(s, "Valores em R$. Fonte: DRE do período (Jan–Jul/2026).")

# ================================================================== SLIDE 5
s = slide(WHITE)
header(s, "Alerta Contábil", "Saldos Invertidos", 5)
alerts = [
    ("140", "Clientes (grupo)", 1123648.52, "ATIVO", "devedora", "credora",
     "Conta de Ativo com natureza normalmente devedora apresentando saldo credor."),
    ("1494", "Fornecedores (grupo)", 441904.06, "PASSIVO", "credora", "devedora",
     "Conta de Passivo com natureza normalmente credora apresentando saldo devedor."),
]
cw2 = 5.9; x0 = 0.62; y0 = 2.0; ch2 = 3.55
for i, (cod, nome, val, grp, nat, atual, desc) in enumerate(alerts):
    x = x0 + i*(cw2+0.29)
    rect(s, x, y0, cw2, ch2, fill=AMBER_BG, line=ORANGE, lw=1.25, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    rect(s, x, y0, cw2, 0.10, fill=ORANGE)
    txt(s, x+0.4, y0+0.35, cw2-0.8, 0.5, [
        [("CONTA ", 13, GRAY, True, BODY_FONT), (cod, 13, ORANGE, True, BODY_FONT),
         ("  ·  ", 13, GRAY, True, BODY_FONT), (nome, 13, GRAY_TX, True, BODY_FONT)]])
    txt(s, x+0.4, y0+0.85, cw2-0.8, 0.3, [[("Saldo invertido (sintética)", 10, GRAY, False, BODY_FONT)]])
    txt(s, x+0.4, y0+1.18, cw2-0.8, 0.6, [[(brl(val), 30, ORANGE, True, TITLE_FONT)]])
    txt(s, x+0.4, y0+1.95, cw2-0.8, 0.45, [
        [("Grupo: ", 10, GRAY, True, BODY_FONT), (grp, 10, GRAY_TX, True, BODY_FONT),
         ("   |   Natureza esperada: ", 10, GRAY, True, BODY_FONT), (nat, 10, GRAY_TX, True, BODY_FONT),
         ("   |   Apresentada: ", 10, GRAY, True, BODY_FONT), (atual, 10, RED, True, BODY_FONT)]])
    txt(s, x+0.4, y0+2.4, cw2-0.8, 0.9, [[(desc, 11, GRAY_TX, False, BODY_FONT)]], spacing=1.15)
# selo
rect(s, 3.55, 6.18, 6.2, 0.62, fill=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
txt(s, 3.55, 6.18, 6.2, 0.62, [[("⚑  NECESSIDADE DE CONCILIAÇÃO IMEDIATA", 13, WHITE, True, BODY_FONT)]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
footer(s, "140 e 1494 são contas sintéticas (grupo). Saldo invertido é indício que exige conciliação — não constitui, por si, prova de irregularidade.")

# ================================================================== SLIDE 6
s = slide(WHITE)
header(s, "Interpretação", "O que os saldos podem indicar", 6)
# infografico compra vs venda
def flowrow(y, titulo, cor, steps):
    txt(s, 0.62, y, 2.0, 0.6, [[(titulo, 13, cor, True, BODY_FONT)]], anchor=MSO_ANCHOR.MIDDLE)
    fx = 2.6; fw = 3.0; fh = 0.62; gapx = 0.15
    for i, st in enumerate(steps):
        x = fx + i*(fw+gapx)
        sh = MSO_SHAPE.CHEVRON
        c = rect(s, x, y, fw, fh, fill=(cor if i == 0 else GRAY_LT), shape=sh)
        tc = WHITE if i == 0 else GRAY_TX
        txt(s, x+0.2, y, fw-0.4, fh, [[(st, 10.5, tc, True, BODY_FONT)]],
            align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
flowrow(2.1, "COMPRA", BLUE, ["Documento fiscal (NF)", "Lançamento contábil", "Financeiro (pagar)"])
flowrow(3.0, "VENDA", GREEN, ["Documento fiscal (NF)", "Lançamento contábil", "Financeiro (receber)"])
txt(s, 2.6, 3.75, 9.5, 0.4, [[("Quando o ciclo documento → lançamento → financeiro não fecha, o saldo pode inverter.",
    10.5, GRAY, False, BODY_FONT)]])
# possiveis causas
rect(s, 0.62, 4.35, 12.09, 2.35, fill=GRAY_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
rect(s, 0.62, 4.35, 0.09, 2.35, fill=ORANGE)
txt(s, 0.95, 4.55, 11.5, 0.4, [[("POSSÍVEIS CAUSAS A INVESTIGAR (hipóteses, não conclusões)", 12, ORANGE, True, BODY_FONT)]])
causas = [
    "Recebimentos/pagamentos lançados sem o respectivo título de origem.",
    "Baixas em duplicidade ou baixa superior ao valor original do título.",
    "Notas fiscais não escrituradas ou escrituradas em conta indevida.",
    "Adiantamentos de clientes/a fornecedores registrados no grupo errado.",
    "Diferenças de conciliação bancária não tratadas.",
    "Erros de classificação entre contas do mesmo grupo sintético.",
]
half = 3
for i, c in enumerate(causas):
    col, rr = divmod(i, half)
    x = 0.95 + col*5.9; y = 5.0 + rr*0.52
    rect(s, x, y+0.06, 0.12, 0.12, fill=ORANGE, shape=MSO_SHAPE.OVAL)
    txt(s, x+0.28, y, 5.5, 0.5, [[(c, 10.5, GRAY_TX, False, BODY_FONT)]], spacing=1.05)
footer(s, "Hipóteses apresentadas apenas como direcionadores da conciliação. Confirmação depende da análise documental.")

# ================================================================== SLIDE 7
s = slide(WHITE)
header(s, "Diagnóstico", "Principais Pontos de Atenção", 7)
pa = [
    ("PREJUÍZO NO PERÍODO", brl(357139.26).replace("R$", "–R$"),
     "Resultado líquido negativo; estrutura de custos e despesas acima da receita.", RED),
    ("CLIENTE COM SALDO INVERTIDO", brl(1123648.52),
     "Conta 140 (grupo) apresenta saldo credor — natureza contábil incompatível.", ORANGE),
    ("FORNECEDOR COM SALDO INVERTIDO", brl(441904.06),
     "Conta 1494 (grupo) apresenta saldo devedor — natureza contábil incompatível.", ORANGE),
    ("CONCILIAÇÃO NECESSÁRIA", "Prioridade",
     "Saldos exigem conciliação e reprocessamento antes de qualquer decisão gerencial.", BLUE),
]
cw3 = 5.9; ch3 = 1.95
for i, (t_, v, d, col) in enumerate(pa):
    r, c = divmod(i, 2)
    x = 0.62 + c*(cw3+0.29); y = 2.0 + r*(ch3+0.28)
    rect(s, x, y, cw3, ch3, fill=WHITE, line=GRAY_BD, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
    rect(s, x, y, 0.12, ch3, fill=col)
    rect(s, x+0.35, y+0.32, 0.55, 0.55, fill=col, shape=MSO_SHAPE.OVAL)
    txt(s, x+0.35, y+0.34, 0.55, 0.5, [[("!", 22, WHITE, True, TITLE_FONT)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, x+1.1, y+0.28, cw3-1.4, 0.35, [[(t_, 11.5, col, True, BODY_FONT)]])
    txt(s, x+1.1, y+0.62, cw3-1.4, 0.4, [[(v, 19, GRAY_TX, True, TITLE_FONT)]])
    txt(s, x+1.1, y+1.12, cw3-1.4, 0.75, [[(d, 10, GRAY, False, BODY_FONT)]], spacing=1.1)
footer(s, "Pontos de atenção descritos em linguagem técnica, sem juízo de irregularidade. Fonte: DRE e Balancete.")

# ================================================================== SLIDE 8
s = slide(WHITE)
header(s, "Recomendação", "Plano de Ação FELCONT", 8)
etapas = [
    ("1", "Levantar documentação", "Reunir NFs, extratos e títulos do período."),
    ("2", "Conciliar Clientes (140)", "Cruzar títulos a receber com baixas e recebimentos."),
    ("3", "Conciliar Fornecedores (1494)", "Cruzar títulos a pagar com baixas e pagamentos."),
    ("4", "Revisar lançamentos", "Identificar baixas em duplicidade e classificações."),
    ("5", "Conciliação bancária", "Ajustar diferenças de caixa e bancos."),
    ("6", "Corrigir e ajustar", "Efetuar os lançamentos de acerto necessários."),
    ("7", "Reprocessar o resultado", "Reapurar DRE após as correções."),
    ("8", "Validar e reportar", "Confirmar saldos e emitir relatório gerencial."),
]
x0 = 0.62; cw4 = 2.75
gapx = (12.09 - 4*cw4) / 3
lineA = 3.55; lineB = 5.05
rect(s, x0+0.2, lineA, 11.55, 0.03, fill=GRAY_BD)
rect(s, x0+0.2, lineB, 11.55, 0.03, fill=GRAY_BD)
for i, (n, t_, d) in enumerate(etapas):
    row = i // 4
    x = x0 + (i % 4)*(cw4+gapx)
    if row == 0:
        card_y = 1.95; line_y = lineA
    else:
        card_y = 5.35; line_y = lineB
    cx = x + cw4/2
    dot = rect(s, cx-0.27, line_y+0.015-0.27, 0.54, 0.54, fill=BLUE, shape=MSO_SHAPE.OVAL)
    txt(s, cx-0.27, line_y-0.27, 0.54, 0.54, [[(n, 17, WHITE, True, TITLE_FONT)]],
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box_y = card_y if row == 0 else card_y
    rect(s, x, box_y, cw4, 1.15, fill=GRAY_LT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.07)
    rect(s, x, box_y, cw4, 0.07, fill=BLUE)
    txt(s, x+0.22, box_y+0.18, cw4-0.44, 0.4, [[(t_, 11, BLUE, True, BODY_FONT)]], spacing=1.0)
    txt(s, x+0.22, box_y+0.56, cw4-0.44, 0.55, [[(d, 9, GRAY_TX, False, BODY_FONT)]], spacing=1.05)
# mensagem-chave
rect(s, 0.62, 6.75, 12.09, 0.5, fill=BLUE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.4)
txt(s, 0.62, 6.75, 12.09, 0.5, [[("Mensagem-chave:  ", 12, RGBColor(0xCF,0xEF,0xFB), True, BODY_FONT),
    ("o resultado só deve ser lido para decisão APÓS a conciliação e o reprocessamento.", 12, WHITE, False, BODY_FONT)]],
    align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ================================================================== SLIDE 9
s = slide(NAVY)
rect(s, 0, 0, SW, 0.14, fill=BLUE)
s.shapes.add_picture(LOGO_PNG, Inches(0.7), Inches(0.55), width=Inches(2.9))
txt(s, 0.75, 1.9, 11.0, 0.5, [[("CONCLUSÃO / PRÓXIMOS PASSOS", 14, BLUE, True, BODY_FONT)]])
txt(s, 0.75, 2.35, 11.5, 0.8, [[("Do diagnóstico à decisão", 34, WHITE, True, TITLE_FONT)]])
# fluxo 4 etapas
fluxo = [("CONCILIAR", "Cruzar documentos e saldos"),
         ("CORRIGIR", "Ajustar lançamentos"),
         ("VALIDAR", "Reprocessar e conferir"),
         ("ANALISAR", "Decidir com dados confiáveis")]
fw = 2.75; gapx = 0.35; x0 = 0.75; y = 3.6
for i, (t_, d) in enumerate(fluxo):
    x = x0 + i*(fw+gapx)
    rect(s, x, y, fw, 1.35, fill=NAVY2, line=BLUE, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    txt(s, x+0.25, y+0.28, fw-0.5, 0.5, [[(t_, 17, BLUE, True, TITLE_FONT)]])
    txt(s, x+0.25, y+0.78, fw-0.5, 0.5, [[(d, 10, WHITE, False, BODY_FONT)]], spacing=1.05)
    if i < 3:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x+fw+0.02), Inches(y+0.55), Inches(0.32), Inches(0.28))
        ar.fill.solid(); ar.fill.fore_color.rgb = ORANGE; ar.line.fill.background(); ar.shadow.inherit = False
# assinatura / slogan
rect(s, 0.78, 5.55, 1.5, 0.06, fill=ORANGE)
txt(s, 0.75, 5.8, 11.5, 0.6, [[("“Informação contábil transformada em decisão.”", 18, WHITE, True, TITLE_FONT)]])
txt(s, 0.75, 6.55, 11.5, 0.5, [
    [("FELCONT — Contabilidade, Finanças e Auditoria", 12, BLUE, True, BODY_FONT)],
    [("Análise gerencial · Cliente: Ronaldo Donadon · Período Jan–Jul/2026", 10, GRAY, False, BODY_FONT)],
], spacing=1.15)

out = "/app/FELCONT_Analise_Gerencial.pptx"
prs.save(out)
print("OK ->", out)
